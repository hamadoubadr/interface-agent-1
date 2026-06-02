from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd
import pdfplumber
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
try:
    import fitz  # type: ignore[import-not-found]
except Exception:
    fitz = None

from .pipeline import normalize_name


GENERIC_PRODUCT_PARTS = {
    "lunite",
    "unite",
    "lensemble",
    "ensemble",
    "lunitedensemble",
    "unitedensemble",
    "lemetrelineaire",
    "metrelineaire",
    "lemetrecarre",
    "metrecarre",
    "forfait",
}

STOPWORDS = {
    "de",
    "des",
    "du",
    "la",
    "le",
    "les",
    "et",
    "ou",
    "a",
    "au",
    "aux",
    "en",
    "pour",
    "avec",
    "sans",
    "sur",
    "par",
    "dans",
    "lot",
}

SPEC_HINTS = (
    "diam",
    "ø",
    "dn",
    "pn",
    "kw",
    "w",
    "btu",
    "m3/h",
    "l/s",
    "bar",
    "mmce",
    "pvc",
    "ppr",
    "pehd",
    "cuivre",
    "acier",
    "epaisseur",
    "épaisseur",
    "pression",
    "puissance",
    "debit",
    "débit",
    "vrv",
    "vrf",
    "vmc",
    "classe",
    "coupe-feu",
    "isolation",
    "temperature",
    "température",
    "calorifuge",
)

MATERIAL_HINTS = (
    "inox",
    "inoxydable",
    "acier",
    "bronze",
    "laiton",
    "fonte",
    "cuivre",
    "polypropylene",
    "polypropyl",
    "galvanise",
    "galvanisé",
    "elastomere",
    "élastomère",
)

EQUIPMENT_HINTS = (
    "vanne",
    "clapet",
    "filtre",
    "disconnecteur",
    "pompe",
    "moteur",
    "surpresseur",
    "collecteur",
    "nourrice",
    "reservoir",
    "réservoir",
    "gaine",
    "diffuseur",
    "bouche",
    "grille",
    "caisson",
    "volet",
    "extincteur",
    "ria",
    "tube",
    "tuyau",
    "siphon",
    "calorifuge",
    "gainable",
)

GENERIC_SPEC_PATTERNS = (
    r"^fourniture(?:\s+et)?\s+pose\b",
    r"^fourniture\b",
    r"^pose\b",
    r"^installation\b",
    r"^mise en service\b",
    r"^l['’]ensemble(?:\s+comprendra|\s+comporte)?\b",
    r"^y compris\b",
    r"^toutes fournitures et suj[eé]tions\b",
    r"^description des ouvrages\b",
    r"^prix\s*n[°o]?\b",
    r"^ouvrage paye\b",
    r"^le prix comprendra\b",
)

DOCUMENT_NOISE_HINTS = (
    "cahier des prescriptions techniques",
    "devis descriptif des ouvrages",
    "bordereau des prix",
    "detail estimatif",
    "navis property",
)

ACCESSORY_BRAND_HINTS = (
    "support",
    "supportage",
    "collier",
    "fixation",
    "plot anti",
    "anti vibratile",
    "manchette",
    "chassis",
    "rosace",
    "robinet d equerre",
    "disjoncteur",
    "courroie",
)

SEMANTIC_TAG_PATTERNS: dict[str, tuple[str, ...]] = {
    "pipe": (
        "tuyau",
        "tube",
        "tuyaut",
        "canalis",
        "conduite",
        "evacuation",
        "refoulement",
        "condensat",
    ),
    "plastic_pipe": (
        "pvc",
        "ppr",
        "pehd",
        "polyethylene",
        "polypropyl",
        "cuivre",
    ),
    "pump": (
        "pompe",
        "electropompe",
        "surpresseur",
        "relevage",
        "prefiltration",
        "compresseur",
    ),
    "hvac_unit": (
        "vrv",
        "vrf",
        "drv",
        "split",
        "climatiseur",
        "groupe exterieur",
        "unite gainable",
        "unite interieure",
        "unite exterieure",
        "gainable",
    ),
    "air_network": (
        "gaine",
        "grille",
        "bouche",
        "diffuseur",
        "volet",
        "clapet",
        "caisson d air neuf",
        "air neuf",
        "vmc",
        "desenfumage",
    ),
    "air_terminal": (
        "grille",
        "bouche",
        "diffuseur",
        "ventouse",
        "air neuf",
        "rejet",
        "extraction",
    ),
    "damper": (
        "clapet",
        "volet",
        "registre",
        "anti retour",
        "pare flamme",
        "coupe feu",
        "dosage",
    ),
    "water_heater": (
        "chauffe eau",
        "ballon",
        "ecs",
    ),
    "sanitary": (
        "siphon",
        "lavabo",
        "wc",
        "bidet",
        "douche",
        "robinet",
        "sanitaire",
        "puisage",
    ),
    "pool": (
        "piscine",
        "bassin",
        "prise de balai",
        "chlore",
        "redox",
        "filtration",
        "projecteur led",
    ),
    "electrical_control": (
        "telecommande",
        "regulation",
        "bus",
        "cablage",
        "coffret",
        "armoire",
        "transformateur",
        "interrupteur",
        "projecteur",
    ),
    "fire": (
        "extincteur",
        "incendie",
        "ria",
        "poteau incendie",
        "pare flamme",
        "coupe feu",
    ),
}

CONTINUATION_TAILS = (
    "ou",
    "et",
    "de",
    "du",
    "des",
    "avec",
    "pour",
    "sur",
    "vers",
    "muni",
    "munie",
    "volumique",
    "diametre que le",
    "diamètre que le",
)

TECHNICAL_VALUE_RE = re.compile(
    r"\b(?:"
    r"\d+(?:[.,]\d+)?\s*(?:kw|w|mm|cm|m3/h|l/s|bar|mce|pa|db(?:\(a\))?|u|hz|v|a|l|m²|m2|m)"
    r"|dn\s*\d+"
    r"|pn\s*\d+"
    r"|ip\s*\d{2}"
    r"|ie\s*\d"
    r"|r(?:32|410a?)"
    r"|nc\s*\d+"
    r"|cf\s*°?\s*\d"
    r")\b",
    re.IGNORECASE,
)

BRAND_PATTERNS = [
    re.compile(r"(?:de\s+marque|marque)\s*[:\-]?\s*[«\"]?(?P<brands>[^.\n\r;]+?)\s*(?:ou\s+(?:equivalent|équivalent|similaire)|,?\s*comprenant|,?\s*compose|$)", re.IGNORECASE),
    re.compile(r"(?:de\s+type|type)\s*[:\-]?\s*[«\"]?(?P<brands>[^.\n\r;]+?)\s*(?:ou\s+(?:equivalent|équivalent|similaire)|,?\s*comprenant|,?\s*compose|$)", re.IGNORECASE),
    re.compile(r"(?:de\s+chez|chez)\s*[:\-]?\s*[«\"]?(?P<brands>[^.\n\r;]+?)\s*(?:ou\b|,?\s*comprenant|,?\s*compose|$)", re.IGNORECASE),
    re.compile(r"\b(?P<brands>[A-Z][A-Za-z0-9+&./-]{1,}(?:\s*,\s*[A-Z][A-Za-z0-9+&./-]{1,}){0,5})\b\s+(?:ou\s+(?:equivalent|équivalent|similaire))", re.IGNORECASE),
]

BRAND_REJECTS = {
    "PVC",
    "PPR",
    "PEHD",
    "VRV",
    "VRF",
    "VMC",
    "DN",
    "PN",
    "LOT",
    "HT",
    "TTC",
    "CVC",
    "ECS",
    "EF",
    "EU",
    "EP",
    "MARQUE",
    "TYPE",
    "SIMILAIRE",
    "EQUIVALENT",
    "ÉQUIVALENT",
}

BRAND_DESCRIPTOR_TOKENS = {
    "SPECIAL",
    "MODELE",
    "MODEL",
    "GAMME",
    "SERIE",
}

BRAND_GENERIC_TERMS = {
    "REGULATION",
    "COLORIMETRIQUE",
    "ALARME",
    "ALARMES",
    "RENVOI",
    "RENVOIS",
    "PRESSION",
    "DEBIT",
    "PUISSANCE",
    "CAGE",
    "ESCALIER",
    "BUS",
    "OUVERT",
    "ENGLOBANT",
    "UNITE",
    "UNITES",
    "UNITÉS",
    "CABLE",
    "CÂBLE",
    "CONDUCTEUR",
    "FILAIRE",
    "FRIGORIGENE",
    "FRIGORIGÈNE",
}


@dataclass(slots=True)
class PdfPage:
    pdf_name: str
    pdf_path: Path
    page_number: int
    text: str
    normalized_text: str
    lines: list[str]


@dataclass(slots=True)
class MatchedBlock:
    page: PdfPage
    lines: list[str]
    score: float
    matched_code: str


def _normalize_search_text(value: str) -> str:
    return normalize_name(value)


def _split_product_parts(product_name: str) -> list[str]:
    parts = [part.strip() for part in re.split(r"\r?\n|\|", product_name or "") if part.strip()]
    cleaned: list[str] = []
    for part in parts:
        norm = _normalize_search_text(part).replace(" ", "")
        if norm in GENERIC_PRODUCT_PARTS:
            continue
        cleaned.append(part)
    return cleaned


def _tokenize(value: str) -> list[str]:
    tokens = []
    for token in _normalize_search_text(value).split():
        if len(token) < 3 or token in STOPWORDS:
            continue
        tokens.append(token)
    return tokens


def _normalize_ocr_punctuation(value: str) -> str:
    text = value or ""
    replacements = {
        "«": '"',
        "»": '"',
        "“": '"',
        "”": '"',
        "„": '"',
        "‟": '"',
        "Ť": '"',
        "ť": '"',
        "’": "'",
        "‘": "'",
        "‚": "'",
        "–": "-",
        "—": "-",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def _semantic_tags_from_text(*values: str) -> set[str]:
    normalized = _normalize_search_text(" ".join(value for value in values if value))
    tags: set[str] = set()
    if not normalized:
        return tags

    for tag, markers in SEMANTIC_TAG_PATTERNS.items():
        if any(marker in normalized for marker in markers):
            tags.add(tag)
    return tags


def _semantic_match_adjustment(product_parts: list[str], chapter: str, subchapter: str, text: str) -> float:
    product_tags = _semantic_tags_from_text(" ".join(product_parts), chapter, subchapter)
    text_tags = _semantic_tags_from_text(text)
    if not product_tags or not text_tags:
        return 0.0

    score = 0.0
    shared = product_tags.intersection(text_tags)
    score += min(len(shared) * 1.8, 4.5)

    if "pipe" in product_tags and any(tag in text_tags for tag in ("pump", "hvac_unit")) and not text_tags.intersection({"pipe", "plastic_pipe", "sanitary"}):
        score -= 9.0
    if "water_heater" in product_tags and any(tag in text_tags for tag in ("pump", "pool", "pipe", "hvac_unit")) and "water_heater" not in text_tags:
        score -= 8.0
    if "hvac_unit" in product_tags and any(tag in text_tags for tag in ("pump", "pool")) and not text_tags.intersection({"hvac_unit", "air_network"}):
        score -= 8.0
    if "air_network" in product_tags and any(tag in text_tags for tag in ("pump", "pool", "water_heater")) and "air_network" not in text_tags:
        score -= 7.0
    if "air_terminal" in product_tags and "damper" in text_tags and "air_terminal" not in text_tags:
        score -= 8.0
    if "damper" in product_tags and "air_terminal" in text_tags and "damper" not in text_tags:
        score -= 6.0
    if "air_terminal" in product_tags and "fire" in text_tags and "fire" not in product_tags:
        score -= 7.0
    if "sanitary" in product_tags and any(tag in text_tags for tag in ("pump", "hvac_unit")) and not text_tags.intersection({"sanitary", "pipe"}):
        score -= 7.0
    if "pool" in product_tags and any(tag in text_tags for tag in ("hvac_unit", "water_heater")) and "pool" not in text_tags:
        score -= 6.0
    if "electrical_control" in product_tags and "pump" in text_tags and not text_tags.intersection({"electrical_control", "hvac_unit"}):
        score -= 5.0

    if "pipe" in product_tags and "plastic_pipe" in text_tags:
        score += 1.0
    if "hvac_unit" in product_tags and text_tags.intersection({"hvac_unit", "air_network"}):
        score += 1.0
    if "air_terminal" in product_tags and text_tags.intersection({"air_terminal", "air_network"}):
        score += 1.0
    if "water_heater" in product_tags and text_tags.intersection({"water_heater", "sanitary"}):
        score += 1.0

    return score


def _extract_with_pypdf(pdf_path: Path) -> list[str]:
    try:
        reader = PdfReader(str(pdf_path), strict=False)
    except Exception:
        return []

    texts: list[str] = []
    for page in reader.pages:
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            texts.append("")
    return texts


def _extract_with_pymupdf(pdf_path: Path) -> list[str]:
    if fitz is None:
        return []

    texts: list[str] = []
    try:
        with fitz.open(pdf_path) as document:
            for page in document:
                try:
                    texts.append(page.get_text("text", sort=True) or "")
                except Exception:
                    texts.append("")
    except Exception:
        return []
    return texts


def _extract_with_pdfplumber(pdf_path: Path) -> list[str]:
    texts: list[str] = []
    try:
        with pdfplumber.open(str(pdf_path)) as pdf:
            for page in pdf.pages:
                try:
                    texts.append(page.extract_text() or "")
                except Exception:
                    texts.append("")
    except Exception:
        return []
    return texts


def _extract_with_binary_fallback(pdf_path: Path) -> list[str]:
    try:
        raw = pdf_path.read_bytes()
    except Exception:
        return []

    text = raw.decode("latin-1", errors="ignore")
    chunks = re.findall(r"[A-Za-zÀ-ÿ0-9][A-Za-zÀ-ÿ0-9 ,;:()/%+°'\"._=-]{5,}", text)
    cleaned_lines: list[str] = []
    for chunk in chunks:
        line = re.sub(r"\s+", " ", chunk).strip()
        if len(line) < 8:
            continue
        cleaned_lines.append(line)
        if len(cleaned_lines) >= 400:
            break

    if not cleaned_lines:
        return []
    return ["\n".join(cleaned_lines)]


def extract_pdf_pages(pdf_path: Path) -> list[PdfPage]:
    texts = _extract_with_pymupdf(pdf_path)
    if sum(len(text.strip()) for text in texts) < 80:
        texts = _extract_with_pypdf(pdf_path)
    if sum(len(text.strip()) for text in texts) < 80:
        texts = _extract_with_pdfplumber(pdf_path)
    if sum(len(text.strip()) for text in texts) < 80:
        texts = _extract_with_binary_fallback(pdf_path)

    pages: list[PdfPage] = []
    for index, text in enumerate(texts, start=1):
        cleaned = text.replace("\u00a0", " ").replace("\u202f", " ").strip()
        lines = [line.strip() for line in cleaned.splitlines() if line.strip()]
        pages.append(
            PdfPage(
                pdf_name=pdf_path.name,
                pdf_path=pdf_path,
                page_number=index,
                text=cleaned,
                normalized_text=_normalize_search_text(cleaned),
                lines=lines,
            )
        )
    return pages


def build_pdf_corpus(pdf_paths: Iterable[Path]) -> list[PdfPage]:
    pages: list[PdfPage] = []
    for path in pdf_paths:
        try:
            pages.extend(extract_pdf_pages(path))
        except Exception:
            continue
    return pages


def _is_toc_line(line: str) -> bool:
    text = line.strip()
    if not text:
        return False
    return bool(
        re.search(r"\.{5,}\s*\d+\s*$", text) or
        re.search(r"\b\d+(?:\.\d+){1,}\b.*\.{5,}\s*\d+\s*$", text) or
        re.fullmatch(r"page", text, re.IGNORECASE)
    )


def _toc_penalty(page: PdfPage) -> float:
    if not page.lines:
        return 0.0

    toc_lines = sum(1 for line in page.lines if _is_toc_line(line))
    ratio = toc_lines / max(len(page.lines), 1)
    penalty = 0.0
    if toc_lines >= 6 or ratio >= 0.20:
        penalty += 8.0
    if "sommaire" in page.normalized_text or "table des matieres" in page.normalized_text:
        penalty += 4.0
    return penalty


def _price_table_penalty(page: PdfPage) -> float:
    if not page.lines:
        return 0.0

    payment_lines = sum(1 for line in page.lines if _is_payment_or_price_line(line))
    code_lines = sum(1 for line in page.lines if re.search(r"\b\d+(?:\.\d+){2,}\b", line))
    dotted_lines = sum(1 for line in page.lines if re.search(r"\.{8,}", line))

    penalty = 0.0
    if payment_lines >= 3:
        penalty += 8.0
    if code_lines >= 8 and payment_lines >= 2:
        penalty += 6.0
    if dotted_lines >= 6:
        penalty += 4.0
    if "bordereau des prix" in page.normalized_text or "devis quantitatif" in page.normalized_text:
        penalty += 6.0
    return penalty


def _normalize_article_code(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip().upper()


def _is_local_subcode(value: str) -> bool:
    return bool(re.fullmatch(r"[a-z]-\d+(?:[.-][a-z0-9]+)?", str(value or "").strip()))


def _line_contains_code(line: str, code_article: str) -> bool:
    normalized_code = _normalize_article_code(code_article)
    if not normalized_code:
        return False
    compact_code = normalized_code.replace(".", "").replace("-", "").replace("/", "")
    compact_line = re.sub(r"\s+", "", str(line or "")).upper().strip()
    compact_line_no_dots = compact_line.replace(".", "").replace("-", "").replace("/", "")

    if compact_line_no_dots == compact_code:
        return True
    if compact_line_no_dots.startswith(compact_code):
        code_len_in_line = 0
        dots_found = 0
        for char in compact_line:
            if char in ".-/":
                dots_found += 1
                continue
            code_len_in_line += 1
            if code_len_in_line == len(compact_code):
                break
        
        total_len = code_len_in_line + dots_found
        if total_len < len(compact_line):
            next_char = compact_line[total_len]
            if next_char.isdigit():
                return False
        return True

    pattern = rf"(?<![A-Z0-9]){re.escape(normalized_code)}(?=$|[^A-Z0-9.])"
    if bool(re.search(pattern, compact_line)):
        return True
    
    flexible_pattern = normalized_code.replace(".", r"\s*\.?\s*")
    if bool(re.search(rf"(?<![A-Z0-9]){flexible_pattern}(?=$|[^A-Z0-9.])", compact_line)):
        return True

    return False


def _looks_like_other_code_line(line: str, code_article: str) -> bool:
    compact_line = re.sub(r"\s+", "", line).upper()
    if not compact_line:
        return False

    if _line_contains_code(line, code_article):
        return False

    if re.search(r"\b[A-Z]-\d+(?:\.[A-Z0-9]+)?\b", compact_line):
        return True
    if re.search(r"\b\d+(?:\.\d+){1,}\b", compact_line):
        return True
    return False


def _is_payment_or_price_line(line: str) -> bool:
    normalized = _normalize_search_text(line)
    return any(
        marker in normalized
        for marker in (
            "ouvrage paye",
            "au prix",
            "prix n",
            "prix no",
            "prix numero",
            "prix",
        )
    )


def _starts_bullet_like(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return False
    first = stripped[0]
    if first in "-•●▪*":
        return True
    if not first.isalnum():
        return True
    if re.match(r"^\d+\b", stripped):
        return True
    return False


def _is_page_header_line(line: str) -> bool:
    normalized = _normalize_search_text(line)
    if not normalized:
        return False
    if normalized in {"maitre d ouvrage", "projet", "titre", "page"}:
        return True
    header_hits = sum(1 for token in ("maitre d ouvrage", "projet", "titre", "page") if token in normalized)
    if header_hits >= 2:
        return True
    if normalized.startswith("b e t ") or normalized.startswith("bet "):
        return True
    if any(token in normalized for token in ("fax", "email", "e mail", "contact@", "contact @")):
        return True
    if re.search(r"\b(?:avenue|bd|boulevard|secteur|hay|capital|s a au capital)\b", normalized) and re.search(r"\d", normalized):
        return True
    return False


def _is_informative_line(line: str) -> bool:
    if not line.strip() or _is_toc_line(line):
        return False
    if _is_payment_or_price_line(line):
        return False
    if _is_page_header_line(line):
        return False
    normalized = _normalize_search_text(line)
    if any(hint in normalized for hint in DOCUMENT_NOISE_HINTS):
        return False
    if len(normalized) >= 30:
        return True
    return bool(
        re.search(r"de\s+marque|marque|ou\s+(?:equivalent|équivalent|similaire)|\btype\b|réf|ref\s*:", line, re.IGNORECASE)
        or any(hint in normalized for hint in SPEC_HINTS)
    )


def _article_ancestors(code_article: str) -> list[str]:
    normalized_code = _normalize_article_code(code_article)
    if "." not in normalized_code:
        return []

    parts = normalized_code.split(".")
    ancestors: list[str] = []
    while len(parts) > 1 and len(ancestors) < 5:
        parts = parts[:-1]
        ancestor = ".".join(parts)
        if ancestor and ancestor not in ancestors:
            ancestors.append(ancestor)
    return ancestors


def _sanitize_code_block_lines(lines: list[str], matched_code: str) -> list[str]:
    sanitized: list[str] = []
    started = not matched_code
    for line in lines:
        if not started:
            if _line_contains_code(line, matched_code):
                started = True
                sanitized.append(line)
            continue

        if _is_toc_line(line):
            continue
        if _is_payment_or_price_line(line):
            break
        if sanitized and _looks_like_other_code_line(line, matched_code):
            break
        sanitized.append(line)

    if sanitized:
        return sanitized
    return [line for line in lines if not _is_toc_line(line) and not _is_payment_or_price_line(line)]


def _needs_continuation(line: str) -> bool:
    normalized = _normalize_search_text(line)
    if not normalized:
        return False
    if any(normalized.endswith(tail) for tail in CONTINUATION_TAILS):
        return True
    stripped = line.strip()
    if stripped.endswith((",", ";", ":", "(")):
        return True
    if stripped.endswith("."):
        return False
    if re.search(r"\bou\s*$", normalized):
        return True
    return bool(re.search(r"[A-Za-zÀ-ÿ0-9]$", stripped))


def _merge_continuation_lines(lines: list[str]) -> list[str]:
    merged: list[str] = []
    for raw_line in lines:
        line = re.sub(r"\s+", " ", raw_line).strip()
        if not line:
            continue
        if not merged:
            merged.append(line)
            continue

        previous = merged[-1]
        normalized_line = _normalize_search_text(line)
        starts_like_continuation = (
            line[:1].islower()
            or normalized_line.startswith(("ou ", "et ", "de ", "du ", "des ", "avec ", "pour "))
            or re.match(r"^(?:[)\].,:;/-]|equivalent\b|équivalent\b)", normalized_line) is not None
        )
        if _needs_continuation(previous) and starts_like_continuation:
            merged[-1] = f"{previous} {line}".strip()
        else:
            merged.append(line)
    return merged


def _score_page(page: PdfPage, product_parts: list[str], chapter: str, subchapter: str) -> float:
    score = 0.0
    normalized_page = page.normalized_text

    for index, part in enumerate(product_parts):
        normalized_part = _normalize_search_text(part)
        if not normalized_part:
            continue
        if normalized_part in normalized_page:
            score += 6.0 if index == 0 else 3.0
        else:
            tokens = _tokenize(part)
            token_hits = sum(1 for token in tokens if token in normalized_page)
            score += min(token_hits * 0.75, 2.5)

    chapter_tokens = _tokenize(chapter)[:3] + _tokenize(subchapter)[:3]
    score += min(sum(1 for token in chapter_tokens if token in normalized_page) * 0.35, 1.2)
    if re.search(r"de\s+marque|marque|ou\s+(?:equivalent|équivalent|similaire)|\btype\b|réf|ref\s*:", page.text, re.IGNORECASE):
        score += 1.5
    score += _semantic_match_adjustment(product_parts, chapter, subchapter, page.text[:2500])
    score += _score_block_context(page.text[:2500], product_parts, chapter, subchapter)
    return score - _toc_penalty(page) - _price_table_penalty(page)


def _score_line(line: str, product_parts: list[str], chapter: str, subchapter: str) -> float:
    if not line.strip() or _is_toc_line(line):
        return -5.0

    normalized_line = _normalize_search_text(line)
    score = 0.0

    for index, part in enumerate(product_parts):
        normalized_part = _normalize_search_text(part)
        if normalized_part and normalized_part in normalized_line:
            score += 8.0 if index == 0 else 4.0
        else:
            tokens = _tokenize(part)
            hits = sum(1 for token in tokens if token in normalized_line)
            score += min(hits * (1.2 if index == 0 else 0.7), 3.0)

    context_tokens = _tokenize(chapter)[:3] + _tokenize(subchapter)[:3]
    score += min(sum(1 for token in context_tokens if token in normalized_line) * 0.4, 1.2)
    if re.search(r"de\s+marque|marque|ou\s+(?:equivalent|équivalent|similaire)|\btype\b|réf|ref\s*:", line, re.IGNORECASE):
        score += 2.5

    return score


def _score_block_context(text: str, product_parts: list[str], chapter: str, subchapter: str) -> float:
    normalized_text = _normalize_search_text(text)
    if not normalized_text:
        return -2.0

    score = 0.0
    product_tokens = _tokenize(" ".join(product_parts[:2]))[:6]
    subchapter_tokens = _tokenize(subchapter)[:5]
    chapter_tokens = _tokenize(chapter)[:4]

    product_hits = sum(1 for token in product_tokens if token in normalized_text)
    subchapter_hits = sum(1 for token in subchapter_tokens if token in normalized_text)
    chapter_hits = sum(1 for token in chapter_tokens if token in normalized_text)

    score += min(product_hits * 1.2, 4.0)
    score += min(subchapter_hits * 1.4, 4.5)
    score += min(chapter_hits * 0.5, 1.5)

    if subchapter_tokens and subchapter_hits == 0:
        score -= 8.0
    elif product_tokens and product_hits == 0:
        score -= 3.0

    return score


def _extract_block_by_code(page: PdfPage, code_article: str, product_parts: list[str], chapter: str, subchapter: str) -> MatchedBlock | None:
    if not code_article or not page.lines:
        return None

    best_lines: list[str] = []
    best_score = float("-inf")
    ancestors = _article_ancestors(code_article)
    for idx, line in enumerate(page.lines):
        if _is_toc_line(line):
            continue
        if not _line_contains_code(line, code_article):
            continue

        start = idx
        backward_context: list[str] = []
        for lookback in range(idx - 1, max(-1, idx - 7), -1):
            if lookback < 0:
                break
            candidate = page.lines[lookback]
            if _is_toc_line(candidate) or _is_payment_or_price_line(candidate):
                break
            if any(_line_contains_code(candidate, ancestor) for ancestor in ancestors):
                backward_context.append(candidate)
                continue
            if _looks_like_other_code_line(candidate, code_article):
                break
            candidate_score = _score_line(candidate, product_parts, chapter, subchapter)
            if _is_informative_line(candidate) or candidate_score >= 1.0:
                backward_context.append(candidate)
                continue
            if backward_context:
                break

        end = idx + 1
        while end < len(page.lines) and end <= idx + 30:
            candidate = page.lines[end]
            if _looks_like_other_code_line(candidate, code_article):
                break
            if _is_payment_or_price_line(candidate):
                break
            end += 1

        window = _sanitize_code_block_lines(list(reversed(backward_context)) + page.lines[start:end], code_article)
        if not window:
            continue

        score = 10.0
        informative = [candidate for candidate in window if _is_informative_line(candidate)]
        score += min(len(informative) * 1.4, 8.0)
        if any(re.search(r"de\s+marque|marque|ou\s+(?:equivalent|équivalent|similaire)|\btype\b|réf|ref\s*:", candidate, re.IGNORECASE) for candidate in window):
            score += 4.0

        for candidate in window:
            score += max(0.0, _score_line(candidate, product_parts, chapter, subchapter))

        score += _score_block_context("\n".join(window), product_parts, chapter, subchapter)
        score -= _toc_penalty(page)
        price_penalty = _price_table_penalty(page)
        if price_penalty > 0.0:
            has_brand_signal = any(
                re.search(r"de\s+marque|marque|de\s+chez|chez|ou\s+(?:equivalent|équivalent|similaire)|\btype\b|réf|ref\s*:", candidate, re.IGNORECASE)
                for candidate in window
            )
            has_exact_code = any(_line_contains_code(candidate, code_article) for candidate in window)
            if has_brand_signal:
                price_penalty *= 0.35
            elif has_exact_code:
                price_penalty *= 0.5
        score -= price_penalty
        if score > best_score:
            best_score = score
            best_lines = window

    if not best_lines:
        return None

    return MatchedBlock(page=page, lines=best_lines, score=best_score, matched_code=code_article)


def _invites_following_bullets(line: str) -> bool:
    normalized = _normalize_search_text(line)
    if not normalized:
        return False
    if _needs_continuation(line):
        return True
    if line.strip().endswith("."):
        return False
    return any(
        marker in normalized
        for marker in (
            "comprenant",
            "comprend",
            "comportant",
            "comporte",
            "compose de",
            "composee de",
            "composees de",
            "composition",
            "caracteristiques",
            "comme suit",
        )
    )


def _extend_block_with_next_page(
    corpus: list[PdfPage],
    block: MatchedBlock,
    product_parts: list[str],
    chapter: str,
    subchapter: str,
) -> MatchedBlock:
    try:
        page_index = corpus.index(block.page)
    except ValueError:
        return block

    if page_index + 1 >= len(corpus):
        return block

    next_page = corpus[page_index + 1]
    if next_page.pdf_path != block.page.pdf_path:
        return block

    logical_lines = _merge_continuation_lines(block.lines)
    meaningful_lines = [line for line in logical_lines if _is_informative_line(line)]
    if not meaningful_lines:
        return block
    if len(block.lines) < 2:
        return block

    last_meaningful = meaningful_lines[-1]
    candidate_preview: list[str] = []
    for line in next_page.lines[:25]:
        if _is_toc_line(line):
            continue
        if _is_page_header_line(line):
            continue
        if any(hint in _normalize_search_text(line) for hint in DOCUMENT_NOISE_HINTS):
            continue
        if _looks_like_other_code_line(line, block.matched_code):
            break
        if _is_payment_or_price_line(line):
            break
        candidate_preview.append(line)

    first_bullet_index = next(
        (index for index, line in enumerate(candidate_preview) if line.strip().startswith(("ï€­", "-", "â€¢", "â—", "â–ª"))),
        -1,
    )
    if first_bullet_index > 0:
        candidate_preview = candidate_preview[first_bullet_index:]

    first_preview = candidate_preview[0].strip() if candidate_preview else ""
    starts_with_bullet = first_preview.startswith(("", "-", "•", "●", "▪"))
    smart_bullet_index = next((index for index, line in enumerate(candidate_preview) if _starts_bullet_like(line)), -1)
    if smart_bullet_index > 0:
        candidate_preview = candidate_preview[smart_bullet_index:]
        first_preview = candidate_preview[0].strip() if candidate_preview else ""
    starts_with_bullet = _starts_bullet_like(first_preview)
    heading_only_block = len(meaningful_lines) == 1 and bool(re.search(r"prix\s*n", meaningful_lines[0], re.IGNORECASE))
    if not _invites_following_bullets(last_meaningful) and not (heading_only_block and starts_with_bullet):
        return block

    appended: list[str] = []
    for line in candidate_preview:
        if _is_toc_line(line):
            continue
        if _is_page_header_line(line):
            continue
        if any(hint in _normalize_search_text(line) for hint in DOCUMENT_NOISE_HINTS):
            continue
        if _looks_like_other_code_line(line, block.matched_code):
            break
        if _is_payment_or_price_line(line):
            break
        appended.append(line)
        if len(appended) >= 6 and not _is_informative_line(line):
            break

    preserved_appended: list[str] = []
    for line in appended:
        if _is_informative_line(line):
            preserved_appended.append(line)
            continue
        if preserved_appended and _needs_continuation(preserved_appended[-1]):
            preserved_appended.append(line)
    appended = preserved_appended
    if not appended:
        return block
    appended_score = _score_block_context("\n".join(appended), product_parts, chapter, subchapter)
    if appended_score < 0.0 and not starts_with_bullet:
        return block

    merged_lines = block.lines + [line for line in appended if line not in block.lines]
    return MatchedBlock(
        page=block.page,
        lines=merged_lines,
        score=block.score + min(len(appended) * 0.4, 2.0),
        matched_code=block.matched_code,
    )


def _find_best_code_block(corpus: list[PdfPage], code_article: str, product_parts: list[str], chapter: str, subchapter: str) -> MatchedBlock | None:
    exact_best: MatchedBlock | None = None
    for page in corpus:
        block = _extract_block_by_code(page, code_article, product_parts, chapter, subchapter)
        if block is None:
            continue
        if _price_table_penalty(page) <= 0.0:
            block.score += 15.0
        block = _extend_block_with_next_page(corpus, block, product_parts, chapter, subchapter)
        if exact_best is None or block.score > exact_best.score:
            exact_best = block

    ancestor_best: MatchedBlock | None = None
    if exact_best is None or exact_best.score < 12.0:
        for depth, ancestor_code in enumerate(_article_ancestors(code_article), start=1):
            for page in corpus:
                block = _extract_block_by_code(page, ancestor_code, product_parts, chapter, subchapter)
                if block is None:
                    continue
                block = _extend_block_with_next_page(corpus, block, product_parts, chapter, subchapter)
                adjusted_score = block.score - (depth * 3.0)
                if _price_table_penalty(block.page) <= 0.0:
                    adjusted_score += 5.0
                
                candidate = MatchedBlock(
                    page=block.page,
                    lines=block.lines,
                    score=adjusted_score,
                    matched_code=block.matched_code,
                )
                if ancestor_best is None or candidate.score > ancestor_best.score:
                    ancestor_best = candidate

    if exact_best and exact_best.score >= 10.0:
        return exact_best
    
    if ancestor_best and ancestor_best.score >= 8.0:
        if exact_best is None or ancestor_best.score >= exact_best.score:
            return ancestor_best
            
    return exact_best or ancestor_best


def _is_financial_noise(value: str) -> bool:
    normalized = _normalize_search_text(value or "")
    if not normalized:
        return False
    if "tva" in normalized or normalized.startswith("t v a"):
        return True
    if normalized.startswith(("total", "sous total", "montant", "net a payer", "a payer")):
        return True
    if any(token in normalized for token in ("ttc", "ht", "remise", "rabais")):
        return True
    return False


def _lines_around_match(page: PdfPage, product_parts: list[str], chapter: str, subchapter: str) -> list[str]:
    lines = page.lines
    if not lines:
        return []

    best_index = -1
    best_score = float("-inf")
    for idx, line in enumerate(lines):
        score = _score_line(line, product_parts, chapter, subchapter)
        if score > best_score:
            best_score = score
            best_index = idx

    if best_index < 0 or best_score < 1.5:
        return [line for line in lines[: min(5, len(lines))] if not _is_toc_line(line)]

    start = max(0, best_index - 2)
    end = min(len(lines), best_index + 4)
    window = [line for line in lines[start:end] if not _is_toc_line(line)]

    if end < len(lines):
        next_line = lines[end]
        if re.search(r"de\s+marque|marque|ou\s+(?:equivalent|équivalent|similaire)|\btype\b|réf|ref\s*:", next_line, re.IGNORECASE):
            window.append(next_line)

    return window


def _normalize_brand_candidate(value: str) -> str:
    text = value.strip(" :;-.,'\"“”«»")
    text = re.sub(r"\s+", " ", text)
    text = re.split(r"\br[ée]f\s*:?", text, flags=re.IGNORECASE)[0].strip()
    text = re.sub(r"^(?:de\s+)?(?:marque|type)\s+", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+ou\s+(?:equivalent|équivalent|similaire)\b.*$", "", text, flags=re.IGNORECASE)
    text = text.strip(" :;-.,'\"“”«»")
    if not text:
        return ""
    normalized_probe = _normalize_search_text(text)
    if normalized_probe in {"equivalent", "similaire"} or "quivalent" in normalized_probe:
        return ""
    text = text.upper()
    text = text.replace("SAULER & PALAU", "SOLER & PALAU")
    if re.fullmatch(r"R(?:32|410A?)", text, re.IGNORECASE):
        return ""
    if any(marker in text for marker in ("U1000", "RO2V", "CR1")):
        return ""
    tokens = text.split()
    for index, token in enumerate(tokens[1:], start=1):
        if token in BRAND_DESCRIPTOR_TOKENS:
            text = " ".join(tokens[:index]).strip()
            tokens = text.split()
            break
    if any(token in BRAND_GENERIC_TERMS for token in tokens):
        return ""
    if len(tokens) > 4 and "&" not in text and "-" not in text and "/" not in text:
        return ""
    return text.strip()


def _split_brand_candidates(value: str) -> list[str]:
    raw_parts = re.split(r",|/|\bou\b|\bet\b", value, flags=re.IGNORECASE)
    brands: list[str] = []
    for raw in raw_parts:
        candidate = _normalize_brand_candidate(raw)
        if not candidate:
            continue
        if len(candidate) < 2:
            continue
        if candidate.upper() in BRAND_REJECTS:
            continue
        normalized_candidate = _normalize_search_text(candidate)
        if normalized_candidate in {"equivalent", "similaire"} or "quivalent" in normalized_candidate:
            continue
        if re.fullmatch(r"(?:equivalent|équivalent|similaire)", candidate, re.IGNORECASE):
            continue
        if candidate not in brands:
            brands.append(candidate)
    return brands


def _extract_brand(text: str) -> str:
    priority_lines = [
        line for line in text.splitlines()
        if _is_informative_line(line) and not any(hint in _normalize_search_text(line) for hint in DOCUMENT_NOISE_HINTS)
    ][:4]
    search_spaces = ["\n".join(priority_lines)] if priority_lines else [text]
    if not priority_lines or len(priority_lines) < 2:
        search_spaces.append(text)

    primary_found: list[str] = []
    secondary_found: list[str] = []
    for search_text in search_spaces:
        for index, pattern in enumerate(BRAND_PATTERNS):
            match = pattern.search(search_text)
            if not match:
                continue
            brands_group = match.groupdict().get("brands") or match.group(1)
            for candidate in _split_brand_candidates(brands_group):
                if index == 0:
                    if candidate not in primary_found:
                        primary_found.append(candidate)
                else:
                    if candidate not in secondary_found:
                        secondary_found.append(candidate)
        if primary_found or secondary_found:
            break
    if not primary_found and not secondary_found:
        normalized_text = _normalize_search_text(text)
        fallback_patterns = (
            re.compile(
                r"(?:de marque|marque|de type|type|de chez|chez)\s+"
                r"(?P<brands>[a-z0-9+&./-]+(?:\s+[a-z0-9+&./-]+){0,3})"
                r"\s+ou\s+(?:equivalent|similaire)",
                re.IGNORECASE,
            ),
            re.compile(
                r"(?:de marque|marque|de type|type|de chez|chez)\s+"
                r"(?P<brands>[a-z0-9+&./-]+(?:\s+[a-z0-9+&./-]+){0,3})"
                r"\s+(?:comprenant|compose|avec)\b",
                re.IGNORECASE,
            ),
        )
        for pattern in fallback_patterns:
            match = pattern.search(normalized_text)
            if not match:
                continue
            brands_group = match.group("brands")
            for candidate in _split_brand_candidates(brands_group):
                if candidate not in secondary_found:
                    secondary_found.append(candidate)
            if secondary_found:
                break
    found = primary_found or secondary_found
    return " | ".join(found[:4])


def _remove_code_prefix(text: str, code_article: str) -> str:
    if not text:
        return ""
    cleaned = text.strip()
    normalized_code = _normalize_article_code(code_article)
    if normalized_code:
        cleaned = re.sub(rf"^\s*{re.escape(normalized_code)}\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^\s*\d+(?:\.\d+){1,}\s*", "", cleaned, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", cleaned).strip(" -:;,.")


def _sanitize_spec_fragment(fragment: str, code_article: str, brand: str) -> str:
    cleaned = _remove_code_prefix(fragment, code_article)
    if not cleaned:
        return ""

    cleaned = cleaned.replace("", " ").replace("\uf0b7", " ").replace("•", " ").replace("●", " ").replace("▪", " ")
    cleaned = cleaned.replace("SAULER & PALAU", "SOLER & PALAU")
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")
    cleaned = re.sub(r"^(?:de\s+marque|marque|de\s+type|type)\s*[:\-]?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(
        r"^(?:fourniture(?:\s*,?\s*pose)?(?:\s+et\s+raccordement)?(?:\s+installation)?(?:\s+et\s+mise\s+en\s+service)?(?:\s+d['’]un[e]?)?\s*)",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(
        r"^(?:et\s+)?pose(?:\s+et\s+raccordement)?(?:\s+d['’]un[e]?)?\s*",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(
        r"^(?:et\s+)?mise\s+en\s+(?:oeuvre|œuvre|.{0,3}uvre)(?:\s+d['’]un[e]?)?\s*",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(
        r"^(?:est\s+mise\s+en\s+service\s+d['’]un[e]?\s*)",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(r"^d['’]?\s*un[e]?\s+", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^l['’]ensemble(?:\s+comprendra|\s+comporte)?\s*:?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^de\s+(?=(?:volet|gaine|cartouche|exutoire|ventilateur|grille|siphon|tube|tuyau|unite|unité|pompe|surpresseur)\b)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")
    normalized_cleaned = _normalize_search_text(cleaned)

    if any(hint in normalized_cleaned for hint in DOCUMENT_NOISE_HINTS):
        return ""
    if normalized_cleaned.startswith("contraintes d installation"):
        return ""
    if normalized_cleaned.startswith("selectionnee en fonction des besoins thermiques"):
        return ""

    if brand:
        first_brand = brand.split("|", 1)[0].strip()
        if first_brand:
            cleaned = re.sub(
                rf",?\s*marque\s+{re.escape(first_brand)}\s+(?:ou\s+(?:equivalent|équivalent|similaire))",
                "",
                cleaned,
                flags=re.IGNORECASE,
            ).strip(" -:;,.")
            cleaned = re.sub(
                rf"\b{re.escape(first_brand)}\b\s*(?:ou\s+(?:equivalent|équivalent|similaire))?",
                "",
                cleaned,
                flags=re.IGNORECASE,
            ).strip(" -:;,.")

    cleaned = re.sub(r"\bou\s+(?:equivalent|équivalent|similaire)\b", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bmarque\b", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bde\s+(?=(?:comprenant|compose|avec)\b)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bde\.\s+(?=[A-ZÀ-ÿ])", "", cleaned)
    cleaned = re.sub(r"\bde\.?\s*$", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")

    return cleaned


def _split_spec_fragments(line: str) -> list[str]:
    normalized_line = re.sub(r"[•●▪]+", ";", line)
    normalized_line = re.sub(r"\s*;\s*", "; ", normalized_line)
    parts = re.split(
        r";\s*|"
        r",\s*(?=(?:\d+|un|une|deux|trois|quatre|clapet|vanne|filtre|disconnecteur|pompe|moteur|reservoir|réservoir|crepine|crépine|collecteur|gaine|bouche|grille|tube|tuyau|marque|type)\b)",
        normalized_line,
        flags=re.IGNORECASE,
    )
    return [part.strip() for part in parts if part.strip()]


def _is_heading_like_fragment(fragment: str, product_parts: list[str]) -> bool:
    normalized = _normalize_search_text(fragment)
    if not normalized:
        return True

    first_product = _normalize_search_text(product_parts[0]) if product_parts else ""
    if first_product and normalized == first_product:
        return True

    if re.fullmatch(r"\d+(?:\.\d+)+", normalized):
        return True

    if first_product:
        first_product_tokens = first_product.split()
        fragment_tokens = normalized.split()
        if len(fragment_tokens) >= 4 and fragment_tokens[:4] == first_product_tokens[:4]:
            return True
        if normalized.startswith(first_product) and len(normalized) <= len(first_product) + 12:
            return True

    if len(normalized) <= 10 and not TECHNICAL_VALUE_RE.search(fragment):
        return True

    letters = [char for char in fragment if char.isalpha()]
    uppercase_ratio = (sum(1 for char in letters if char.isupper()) / len(letters)) if letters else 0.0
    if uppercase_ratio >= 0.85 and len(normalized.split()) <= 4 and not TECHNICAL_VALUE_RE.search(fragment):
        return True

    return False


def _is_technical_fragment(fragment: str) -> bool:
    normalized = _normalize_search_text(fragment)
    if not normalized:
        return False

    if any(re.search(pattern, fragment, re.IGNORECASE) for pattern in GENERIC_SPEC_PATTERNS):
        return False

    if TECHNICAL_VALUE_RE.search(fragment):
        return True
    if any(hint in normalized for hint in SPEC_HINTS):
        return True
    if any(hint in normalized for hint in MATERIAL_HINTS):
        return True
    if any(hint in normalized for hint in EQUIPMENT_HINTS):
        return True
    if re.search(r"\b(?:marque|type|equivalent|équivalent|similaire)\b", fragment, re.IGNORECASE):
        return True
    return False


def _score_spec_fragment(fragment: str, product_parts: list[str], brand: str) -> float:
    normalized = _normalize_search_text(fragment)
    score = 0.0

    if TECHNICAL_VALUE_RE.search(fragment):
        score += 2.5
    if any(hint in normalized for hint in SPEC_HINTS):
        score += 2.0
    if any(hint in normalized for hint in MATERIAL_HINTS):
        score += 1.5
    if any(hint in normalized for hint in EQUIPMENT_HINTS):
        score += 1.0
    if re.search(r"\b(?:marque|type|equivalent|équivalent|similaire)\b", fragment, re.IGNORECASE):
        score += 0.5

    first_product = _normalize_search_text(product_parts[0]) if product_parts else ""
    if first_product and normalized == first_product:
        score -= 3.0
    elif first_product and normalized.startswith(first_product) and len(normalized) <= len(first_product) + 12:
        score -= 1.5

    if brand:
        first_brand = _normalize_search_text(brand.split("|", 1)[0])
        if first_brand and first_brand in normalized:
            score += 0.4

    if len(fragment) > 220:
        score -= 1.0

    return score


def _dedupe_spec_fragments(fragments: list[str]) -> list[str]:
    kept: list[str] = []
    normalized_kept: list[str] = []
    for fragment in fragments:
        normalized = _normalize_search_text(fragment)
        if not normalized:
            continue
        if any(normalized == existing or normalized in existing or existing in normalized for existing in normalized_kept):
            continue
        kept.append(fragment)
        normalized_kept.append(normalized)
    return kept


def _extract_specs_compact(lines: list[str], product_parts: list[str], code_article: str, brand: str) -> str:
    candidates: list[tuple[int, float, str]] = []
    prepared_lines = _merge_continuation_lines(lines)

    candidate_index = 0
    for part in product_parts[1:]:
        cleaned_part = _sanitize_spec_fragment(part, code_article, brand)
        if not cleaned_part:
            continue
        if _is_technical_fragment(cleaned_part):
            score = 1.8 + _score_spec_fragment(cleaned_part, product_parts, brand)
            if score >= 1.5:
                candidates.append((candidate_index, score, cleaned_part))
                candidate_index += 1

    for line in prepared_lines:
        for raw_fragment in _split_spec_fragments(line):
            fragment = _sanitize_spec_fragment(raw_fragment, code_article, brand)
            if not fragment or _is_heading_like_fragment(fragment, product_parts):
                continue
            if not _is_technical_fragment(fragment):
                continue
            score = _score_spec_fragment(fragment, product_parts, brand)
            if score < 1.5:
                continue
            candidates.append((candidate_index, score, fragment))
            candidate_index += 1

    ranked = sorted(candidates, key=lambda item: (-item[1], item[0], len(item[2])))
    shortlist = ranked[:12]
    ordered = [fragment for _, _, fragment in sorted(shortlist, key=lambda item: item[0])]
    deduped = _dedupe_spec_fragments(ordered)

    selected: list[str] = []
    total_length = 0
    for fragment in deduped:
        projected = total_length + len(fragment) + (3 if selected else 0)
        if selected and (len(selected) >= 6 or projected > 420):
            break
        selected.append(fragment)
        total_length = projected

    return " | ".join(selected)


def _extract_specs(lines: list[str], product_parts: list[str], brand: str = "") -> str:
    specs: list[str] = []
    prepared_lines = _merge_continuation_lines(lines)
    additional_parts = product_parts[1:]
    for part in additional_parts:
        cleaned_part = _sanitize_spec_fragment(part, "", brand)
        if cleaned_part and cleaned_part not in specs:
            specs.append(cleaned_part)

    for line in prepared_lines:
        cleaned_line = _sanitize_spec_fragment(line, "", brand)
        if not cleaned_line:
            continue
        normalized = _normalize_search_text(cleaned_line)
        if any(hint in normalized for hint in SPEC_HINTS):
            if cleaned_line not in specs:
                specs.append(cleaned_line)
        elif re.search(r"\b\d+(?:[.,]\d+)?\b", cleaned_line) and any(symbol in cleaned_line for symbol in ("Ø", "x", "X", "/", "mm", "DN", "PN", "KW", "kW", "BTU")):
            if cleaned_line not in specs:
                specs.append(cleaned_line)
        if len(specs) >= 5:
            break

    return " | ".join(specs[:5])


def _extract_specs_rich(lines: list[str], product_parts: list[str], brand: str = "") -> str:
    specs: list[str] = []
    prepared_lines = _merge_continuation_lines(lines)
    for part in product_parts[1:]:
        cleaned_part = _sanitize_spec_fragment(part, "", brand)
        if cleaned_part and cleaned_part not in specs:
            specs.append(cleaned_part)

    for line in prepared_lines:
        if not _is_informative_line(line):
            continue
        cleaned_line = _sanitize_spec_fragment(line, "", brand)
        if not cleaned_line:
            continue
        normalized = _normalize_search_text(cleaned_line)
        if any(hint in normalized for hint in SPEC_HINTS) or len(normalized) >= 35:
            if cleaned_line not in specs:
                specs.append(cleaned_line)
        elif re.search(r"\b\d+(?:[.,]\d+)?\b", cleaned_line) and any(symbol in cleaned_line for symbol in ("Ø", "x", "X", "/", "mm", "DN", "PN", "KW", "kW", "BTU")):
            if cleaned_line not in specs:
                specs.append(cleaned_line)
        if len(specs) >= 5:
            break

    return " | ".join(specs[:5])


def _clean_display_product_name(product_name: str) -> str:
    cleaned_parts = _split_product_parts(product_name)
    return " | ".join(cleaned_parts) if cleaned_parts else product_name.strip()


def _format_numeric_value(value: object) -> str:
    if value is None:
        return ""
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return str(value).strip()
    if math.isnan(numeric):
        return ""
    if numeric.is_integer():
        return str(int(numeric))
    return str(numeric).replace(".", ",")


def _collect_repeated_line_signatures(corpus: list[PdfPage]) -> set[str]:
    counts: dict[str, int] = {}
    for page in corpus:
        seen: set[str] = set()
        for line in page.lines:
            normalized = _normalize_search_text(line)
            if len(normalized) < 24:
                continue
            if normalized in seen:
                continue
            seen.add(normalized)
        for normalized in seen:
            counts[normalized] = counts.get(normalized, 0) + 1

    return {normalized for normalized, count in counts.items() if count >= 3}


def _filter_noise_lines(lines: list[str], repeated_signatures: set[str], matched_code: str) -> list[str]:
    filtered: list[str] = []
    for line in lines:
        normalized = _normalize_search_text(line)
        if any(hint in normalized for hint in DOCUMENT_NOISE_HINTS):
            continue
        if _is_page_header_line(line):
            continue
        if (
            normalized in repeated_signatures
            and not (matched_code and _line_contains_code(line, matched_code))
            and not re.search(r"de\s+marque|marque|ou\s+(?:equivalent|Ã©quivalent|similaire)|\btype\b|rÃ©f|ref\s*:", line, re.IGNORECASE)
            and not any(hint in normalized for hint in SPEC_HINTS)
        ):
            continue
        filtered.append(line)

    return filtered or lines


def _is_power_only_product(product_name: str) -> bool:
    return bool(re.fullmatch(r"\s*\d+(?:[.,]\d+)?\s*k\s*w\s*", str(product_name or ""), re.IGNORECASE))


def _build_power_context_specs(record: dict[str, object], matched_lines: list[str], brand: str) -> str:
    parts: list[str] = []
    product_name = str(record.get("product_name") or "").strip()

    for line in matched_lines:
        if _is_payment_or_price_line(line):
            continue
        cleaned = re.sub(r"\.{4,}", " ", line)
        cleaned = re.sub(r"^\d+(?:\.\d+)+\s*", "", cleaned).strip(" -")
        cleaned = re.sub(r"\b\d+(?:\.\d+){3,}\b", "", cleaned)
        cleaned = re.sub(r"\bU\s+\d+\s+-?\s+L['’]unité\b", "", cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r"\bU\s+\d+\s+-?\s+L['’]unite\b", "", cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r"\bU\s+\d+\b", "", cleaned)
        cleaned = re.sub(r"\b\d+(?:[.,]\d+)?\s*kW\b", "", cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r"\s+", " ", cleaned).strip()
        if not cleaned:
            continue
        normalized = _normalize_search_text(cleaned)
        if any(hint in normalized for hint in DOCUMENT_NOISE_HINTS):
            continue
        if normalized.startswith("l unite") or normalized.startswith("le metre"):
            continue
        if re.search(r"\b\d+(?:[.,]\d+)?\b", cleaned) and len(normalized.split()) <= 3:
            continue
        if cleaned not in parts:
            parts.append(cleaned)

    if product_name and product_name not in parts:
        parts.append(product_name)

    quantity = _format_numeric_value(record.get("quantite"))
    unite = str(record.get("unite") or "").strip()
    if quantity:
        quantity_label = f"Quantité {quantity}"
        if unite:
            quantity_label = f"{quantity_label} {unite}"
        if quantity_label not in parts:
            parts.append(quantity_label)

    if brand:
        brand_label = f"Marque {brand}"
        if brand_label not in parts:
            parts.append(brand_label)

    filtered_parts: list[str] = []
    for part in parts:
        normalized = _normalize_search_text(part)
        if not normalized:
            continue
        if normalized in {"climatisation", "systeme vrf", "unite exterieure", "unite interieure"} or len(filtered_parts) < 3:
            if part not in filtered_parts:
                filtered_parts.append(part)
            continue
        if normalized.startswith("quantite") or normalized.startswith("marque"):
            if part not in filtered_parts:
                filtered_parts.append(part)

    return " | ".join(filtered_parts[:6])


def _extract_price_table_brand(corpus: list[PdfPage], page: PdfPage) -> str:
    texts = [page.text]
    try:
        page_index = corpus.index(page)
    except ValueError:
        page_index = -1

    if page_index >= 0:
        for offset in (1, 2):
            neighbor_index = page_index + offset
            if neighbor_index >= len(corpus):
                break
            neighbor = corpus[neighbor_index]
            if neighbor.pdf_path != page.pdf_path:
                break
            texts.append(neighbor.text[:2000])

    context = "\n".join(texts)
    patterns = (
        re.compile(r"prix\s+unitaire\s*\(([^)]+)\)", re.IGNORECASE),
        re.compile(r"montant\s+total\s*\(([^)]+)\)", re.IGNORECASE),
        re.compile(r"\(([A-Za-z][A-Za-z0-9+&./-]{2,})\)\s*en\s+d[ha-]*ht\b", re.IGNORECASE),
        re.compile(r"\b([A-Za-z][A-Za-z0-9+&./-]{2,})\s+en\s+dh\s+ht\b", re.IGNORECASE),
        re.compile(r"prix\s+unitaire\s+([A-Za-z][A-Za-z0-9+&./-]{2,})\s+en\s+dh\s+ht\b", re.IGNORECASE),
    )
    for pattern in patterns:
        for match in pattern.finditer(context):
            candidate = _normalize_brand_candidate(match.group(1))
            if not candidate:
                continue
            if candidate.upper() in BRAND_REJECTS:
                continue
            if len(candidate) < 3:
                continue
            return candidate.upper()
    return ""


def _find_power_table_block(corpus: list[PdfPage], code_article: str, chapter: str, subchapter: str) -> MatchedBlock | None:
    target_subchapter = _normalize_search_text(subchapter)
    ancestors = _article_ancestors(code_article)
    best_candidate: MatchedBlock | None = None

    for page in corpus:
        normalized_page = page.normalized_text
        if _toc_penalty(page) > 0.0:
            continue
        if _price_table_penalty(page) <= 0.0:
            continue
        if not any(marker in normalized_page for marker in ("prix unitaire", "quantite", "montant total", "bordereau des prix")):
            continue
        for idx, line in enumerate(page.lines):
            if _is_toc_line(line):
                continue
            if not _line_contains_code(line, code_article):
                continue

            collected: list[str] = []
            for lookback in range(max(0, idx - 5), idx):
                candidate = page.lines[lookback]
                if any(_line_contains_code(candidate, ancestor) for ancestor in ancestors[:3]):
                    if candidate not in collected:
                        collected.append(candidate)
                elif target_subchapter and target_subchapter in _normalize_search_text(candidate):
                    if candidate not in collected:
                        collected.append(candidate)

            if not collected:
                for lookback in range(max(0, idx - 3), idx):
                    candidate = page.lines[lookback]
                    if target_subchapter and target_subchapter in _normalize_search_text(candidate):
                        collected.append(candidate)

            collected.append(line)
            if idx + 1 < len(page.lines) and not _is_payment_or_price_line(page.lines[idx + 1]):
                collected.append(page.lines[idx + 1])

            score = 1.2
            if "vrf" in normalized_page or "vrv" in normalized_page:
                score += 0.9
            if target_subchapter and target_subchapter in normalized_page:
                score += 0.6
            if chapter and _normalize_search_text(chapter) in normalized_page:
                score += 0.4
            if best_candidate is None or score > best_candidate.score:
                best_candidate = MatchedBlock(page=page, lines=collected, score=score, matched_code=code_article)

    return best_candidate


def _article_ancestors(code_article: str) -> list[str]:
    normalized_code = _normalize_article_code(code_article)
    if "." not in normalized_code:
        return []

    parts = normalized_code.split(".")
    ancestors: list[str] = []
    while len(parts) > 1 and len(ancestors) < 4:
        parts = parts[:-1]
        ancestor = ".".join(parts).strip(".")
        if ancestor and ancestor not in ancestors:
            ancestors.append(ancestor)
    return ancestors


def _score_page(page: PdfPage, product_parts: list[str], chapter: str, subchapter: str) -> float:
    score = 0.0
    normalized_page = page.normalized_text

    for index, part in enumerate(product_parts):
        normalized_part = _normalize_search_text(part)
        if not normalized_part:
            continue
        if normalized_part in normalized_page:
            score += 6.0 if index == 0 else 3.0
        else:
            tokens = _tokenize(part)
            token_hits = sum(1 for token in tokens if token in normalized_page)
            score += min(token_hits * 0.75, 2.5)

    chapter_tokens = _tokenize(chapter)[:3] + _tokenize(subchapter)[:3]
    score += min(sum(1 for token in chapter_tokens if token in normalized_page) * 0.35, 1.2)
    if re.search(r"de\s+marque|marque|ou\s+(?:equivalent|Ã©quivalent|similaire)|\btype\b|rÃ©f|ref\s*:", page.text, re.IGNORECASE):
        score += 1.5
    score += _semantic_match_adjustment(product_parts, chapter, subchapter, page.text[:2500])
    score += _score_block_context(page.text[:2500], product_parts, chapter, subchapter)
    return score - _toc_penalty(page) - _price_table_penalty(page)


def _score_line(line: str, product_parts: list[str], chapter: str, subchapter: str) -> float:
    if not line.strip() or _is_toc_line(line):
        return -5.0

    normalized_line = _normalize_search_text(line)
    score = 0.0

    for index, part in enumerate(product_parts):
        normalized_part = _normalize_search_text(part)
        if normalized_part and normalized_part in normalized_line:
            score += 8.0 if index == 0 else 4.0
        else:
            tokens = _tokenize(part)
            hits = sum(1 for token in tokens if token in normalized_line)
            score += min(hits * (1.2 if index == 0 else 0.7), 3.0)

    context_tokens = _tokenize(chapter)[:3] + _tokenize(subchapter)[:3]
    score += min(sum(1 for token in context_tokens if token in normalized_line) * 0.4, 1.2)
    if re.search(r"de\s+marque|marque|ou\s+(?:equivalent|Ã©quivalent|similaire)|\btype\b|rÃ©f|ref\s*:", line, re.IGNORECASE):
        score += 2.5
    score += _semantic_match_adjustment(product_parts, chapter, subchapter, line)
    return score


def _score_block_context(text: str, product_parts: list[str], chapter: str, subchapter: str) -> float:
    normalized_text = _normalize_search_text(text)
    if not normalized_text:
        return -2.0

    score = 0.0
    product_tokens = _tokenize(" ".join(product_parts[:2]))[:6]
    subchapter_tokens = _tokenize(subchapter)[:5]
    chapter_tokens = _tokenize(chapter)[:4]

    product_hits = sum(1 for token in product_tokens if token in normalized_text)
    subchapter_hits = sum(1 for token in subchapter_tokens if token in normalized_text)
    chapter_hits = sum(1 for token in chapter_tokens if token in normalized_text)

    score += min(product_hits * 1.2, 4.0)
    score += min(subchapter_hits * 1.4, 4.5)
    score += min(chapter_hits * 0.5, 1.5)

    if subchapter_tokens and subchapter_hits == 0:
        score -= 8.0
    elif product_tokens and product_hits == 0:
        score -= 3.0

    score += _semantic_match_adjustment(product_parts, chapter, subchapter, text)
    return score


def _normalize_brand_candidate(value: str) -> str:
    text = _normalize_ocr_punctuation(value).strip(" :;-.,'\"â€œâ€Â«Â»")
    text = re.sub(r"\s+", " ", text)
    text = re.split(r"\br[Ã©e]f\s*:?", text, flags=re.IGNORECASE)[0].strip()
    text = re.sub(r"^(?:de\s+)?(?:marque|type)\s+", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+ou\s+(?:equivalent|Ã©quivalent|similaire)\b.*$", "", text, flags=re.IGNORECASE)
    text = text.strip(" :;-.,'\"â€œâ€Â«Â»")
    if not text:
        return ""
    normalized_probe = _normalize_search_text(text)
    if normalized_probe in {"equivalent", "similaire"} or "quivalent" in normalized_probe:
        return ""
    text = text.upper()
    text = text.replace("SAULER & PALAU", "SOLER & PALAU")
    if re.fullmatch(r"R(?:32|410A?)", text, re.IGNORECASE):
        return ""
    if any(marker in text for marker in ("U1000", "RO2V", "CR1")):
        return ""
    tokens = text.split()
    for index, token in enumerate(tokens[1:], start=1):
        if token in BRAND_DESCRIPTOR_TOKENS:
            text = " ".join(tokens[:index]).strip()
            tokens = text.split()
            break
    if any(token in BRAND_GENERIC_TERMS for token in tokens):
        return ""
    if len(tokens) > 4 and "&" not in text and "-" not in text and "/" not in text:
        return ""
    return text.strip()


def _score_spec_fragment(fragment: str, product_parts: list[str], brand: str) -> float:
    normalized = _normalize_search_text(fragment)
    score = 0.0

    if TECHNICAL_VALUE_RE.search(fragment):
        score += 2.5
    if any(hint in normalized for hint in SPEC_HINTS):
        score += 2.0
    if any(hint in normalized for hint in MATERIAL_HINTS):
        score += 1.5
    if any(hint in normalized for hint in EQUIPMENT_HINTS):
        score += 1.0
    if re.search(r"\b(?:marque|type|equivalent|Ã©quivalent|similaire)\b", fragment, re.IGNORECASE):
        score += 0.5

    first_product = _normalize_search_text(product_parts[0]) if product_parts else ""
    if first_product and normalized == first_product:
        score -= 3.0
    elif first_product and normalized.startswith(first_product) and len(normalized) <= len(first_product) + 12:
        score -= 1.5

    if brand:
        first_brand = _normalize_search_text(brand.split("|", 1)[0])
        if first_brand and first_brand in normalized:
            score += 0.4

    score += _semantic_match_adjustment(product_parts, "", "", fragment)

    if len(fragment) > 220:
        score -= 1.0

    return score


def _score_brand_line(line: str, product_parts: list[str], chapter: str, subchapter: str) -> float:
    normalized = _normalize_search_text(line)
    score = _semantic_match_adjustment(product_parts, chapter, subchapter, line)
    if "marque" in normalized or "de type" in normalized or "de chez" in normalized or "chez " in normalized:
        score += 2.5
    if any(hint in normalized for hint in ACCESSORY_BRAND_HINTS):
        score -= 2.5
    return score


def _extract_brand(text: str, product_parts: list[str] | None = None, chapter: str = "", subchapter: str = "") -> str:
    product_parts = product_parts or []
    cleaned_text = _normalize_ocr_punctuation(text)
    informative_lines = [
        _normalize_ocr_punctuation(line)
        for line in cleaned_text.splitlines()
        if _is_informative_line(line) and not any(hint in _normalize_search_text(line) for hint in DOCUMENT_NOISE_HINTS)
    ]

    scored_candidates: list[tuple[float, int, str]] = []
    for line_index, line in enumerate(informative_lines[:12]):
        for pattern_index, pattern in enumerate(BRAND_PATTERNS):
            match = pattern.search(line)
            if not match:
                continue
            brands_group = match.groupdict().get("brands") or match.group(1)
            base_score = _score_brand_line(line, product_parts, chapter, subchapter) + max(0.0, 1.6 - (pattern_index * 0.2))
            if line_index < 4:
                base_score += 1.0
            for candidate in _split_brand_candidates(brands_group):
                scored_candidates.append((base_score, line_index, candidate))

    if not scored_candidates:
        normalized_text = _normalize_search_text(cleaned_text)
        fallback_patterns = (
            re.compile(
                r"(?:de marque|marque|de type|type|de chez|chez)\s+"
                r"(?P<brands>[a-z0-9+&./-]+(?:\s+[a-z0-9+&./-]+){0,3})"
                r"\s+ou\s+(?:equivalent|similaire)",
                re.IGNORECASE,
            ),
            re.compile(
                r"(?:de marque|marque|de type|type|de chez|chez)\s+"
                r"(?P<brands>[a-z0-9+&./-]+(?:\s+[a-z0-9+&./-]+){0,3})"
                r"\s+(?:comprenant|compose|avec)\b",
                re.IGNORECASE,
            ),
        )
        for pattern in fallback_patterns:
            match = pattern.search(normalized_text)
            if not match:
                continue
            brands_group = match.group("brands")
            for candidate in _split_brand_candidates(brands_group):
                scored_candidates.append((1.0, 99, candidate))
            if scored_candidates:
                break

    if not scored_candidates:
        return ""

    best_scores: dict[str, tuple[float, int]] = {}
    for score, line_index, candidate in scored_candidates:
        current = best_scores.get(candidate)
        if current is None or score > current[0] or (score == current[0] and line_index < current[1]):
            best_scores[candidate] = (score, line_index)

    ranked = sorted(best_scores.items(), key=lambda item: (-item[1][0], item[1][1], item[0]))
    if not ranked or ranked[0][1][0] < 0.8:
        return ""
    return " | ".join(candidate for candidate, _ in ranked[:4])


def _normalize_brand_candidate(value: str) -> str:
    text = _normalize_ocr_punctuation(value).strip(" :;-.,'\"")
    text = re.sub(r"\s+", " ", text)
    text = re.split(r"\br[eé]f\s*:?", text, flags=re.IGNORECASE)[0].strip()
    text = re.sub(r"^(?:de\s+)?(?:marque|type)\s+", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+ou\s+(?:equivalent|équivalent|similaire)\b.*$", "", text, flags=re.IGNORECASE)
    text = text.strip(" :;-.,'\"")
    if not text:
        return ""

    normalized_probe = _normalize_search_text(text)
    if normalized_probe in {"equivalent", "similaire"} or "quivalent" in normalized_probe:
        return ""

    text = text.upper()
    text = text.replace("SAULER & PALAU", "SOLER & PALAU")
    if re.fullmatch(r"R(?:32|410A?)", text, re.IGNORECASE):
        return ""
    if any(marker in text for marker in ("U1000", "RO2V", "CR1")):
        return ""

    tokens = text.split()
    technical_brand_rejects = {
        "MONOBLOC",
        "SCROLL",
        "INVERTER",
        "VERTICAL",
        "HORIZONTAL",
        "CENTRIFUGE",
        "LINEAIRE",
        "LINÉAIRE",
    }
    if any(token in technical_brand_rejects for token in tokens):
        return ""

    for index, token in enumerate(tokens[1:], start=1):
        if token in BRAND_DESCRIPTOR_TOKENS:
            text = " ".join(tokens[:index]).strip()
            tokens = text.split()
            break

    if any(token in BRAND_GENERIC_TERMS for token in tokens):
        return ""
    if len(tokens) > 4 and "&" not in text and "-" not in text and "/" not in text:
        return ""
    return text.strip()


def _extract_explicit_brand_groups(line: str) -> list[str]:
    cleaned_line = _normalize_ocr_punctuation(line)
    patterns = (
        re.compile(r"(?:de\s+marque|marque|de\s+chez|chez)\s*[:\-]?\s*(?P<brands>.+)$", re.IGNORECASE),
    )
    groups: list[str] = []
    for pattern in patterns:
        match = pattern.search(cleaned_line)
        if not match:
            continue
        brands_group = match.group("brands")
        brands_group = re.split(
            r"\b(?:comprenant|compose|comporte|avec|y\s+compris|caracteristiques?|caract[ée]ristiques?|nota)\b",
            brands_group,
            maxsplit=1,
            flags=re.IGNORECASE,
        )[0]
        brands_group = brands_group.strip().strip(" :;-.,'\"")
        brands_group = re.sub(r"^[\"']+|[\"']+$", "", brands_group).strip()
        if brands_group:
            groups.append(brands_group)
    return groups


def _extract_brand(text: str, product_parts: list[str] | None = None, chapter: str = "", subchapter: str = "") -> str:
    product_parts = product_parts or []
    cleaned_text = _normalize_ocr_punctuation(text)
    informative_lines = [
        _normalize_ocr_punctuation(line)
        for line in cleaned_text.splitlines()
        if _is_informative_line(line) and not any(hint in _normalize_search_text(line) for hint in DOCUMENT_NOISE_HINTS)
    ]

    scored_candidates: list[tuple[float, int, str]] = []
    for line_index, line in enumerate(informative_lines[:12]):
        explicit_groups = _extract_explicit_brand_groups(line)
        if explicit_groups:
            base_score = _score_brand_line(line, product_parts, chapter, subchapter) + 3.0
            if line_index < 4:
                base_score += 1.0
            for brands_group in explicit_groups:
                for candidate in _split_brand_candidates(brands_group):
                    scored_candidates.append((base_score, line_index, candidate))
            continue

        for pattern_index, pattern in enumerate(BRAND_PATTERNS):
            match = pattern.search(line)
            if not match:
                continue
            brands_group = match.groupdict().get("brands") or match.group(1)
            base_score = _score_brand_line(line, product_parts, chapter, subchapter) + max(0.0, 1.6 - (pattern_index * 0.2))
            if line_index < 4:
                base_score += 1.0
            for candidate in _split_brand_candidates(brands_group):
                scored_candidates.append((base_score, line_index, candidate))

    if not scored_candidates:
        normalized_text = _normalize_search_text(cleaned_text)
        fallback_patterns = (
            re.compile(
                r"(?:de marque|marque|de type|type|de chez|chez)\s+"
                r"(?P<brands>[a-z0-9+&./-]+(?:\s+[a-z0-9+&./-]+){0,3})"
                r"\s+ou\s+(?:equivalent|similaire)",
                re.IGNORECASE,
            ),
            re.compile(
                r"(?:de marque|marque|de type|type|de chez|chez)\s+"
                r"(?P<brands>[a-z0-9+&./-]+(?:\s+[a-z0-9+&./-]+){0,3})"
                r"\s+(?:comprenant|compose|avec)\b",
                re.IGNORECASE,
            ),
        )
        for pattern in fallback_patterns:
            match = pattern.search(normalized_text)
            if not match:
                continue
            brands_group = match.group("brands")
            for candidate in _split_brand_candidates(brands_group):
                scored_candidates.append((1.0, 99, candidate))
            if scored_candidates:
                break

    if not scored_candidates:
        return ""

    best_scores: dict[str, tuple[float, int]] = {}
    for score, line_index, candidate in scored_candidates:
        current = best_scores.get(candidate)
        if current is None or score > current[0] or (score == current[0] and line_index < current[1]):
            best_scores[candidate] = (score, line_index)

    ranked = sorted(best_scores.items(), key=lambda item: (-item[1][0], item[1][1], item[0]))
    filtered_ranked: list[tuple[str, tuple[float, int]]] = []
    for candidate, meta in ranked:
        normalized_candidate = _normalize_search_text(candidate)
        is_subphrase = any(
            normalized_candidate
            and normalized_candidate != _normalize_search_text(other_candidate)
            and normalized_candidate in _normalize_search_text(other_candidate).split()
            for other_candidate, _ in ranked
        )
        if not is_subphrase:
            filtered_ranked.append((candidate, meta))

    ranked = filtered_ranked
    if not ranked or ranked[0][1][0] < 0.8:
        return ""
    return " | ".join(candidate for candidate, _ in ranked[:4])


def _split_spec_fragments(line: str) -> list[str]:
    normalized_line = re.sub(r"[â€¢â—ï€­â–ª]+", ";", line)
    normalized_line = re.sub(r"\s*;\s*", "; ", normalized_line)
    parts = re.split(
        r";\s*|"
        r"(?<!\d),\s*(?=(?:\d+|un|une|deux|trois|quatre|clapet|vanne|filtre|disconnecteur|pompe|moteur|reservoir|rÃ©servoir|crepine|crÃ©pine|collecteur|gaine|bouche|grille|tube|tuyau|marque|type)\b)",
        normalized_line,
        flags=re.IGNORECASE,
    )
    return [part.strip() for part in parts if part.strip()]


def _sanitize_spec_fragment(fragment: str, code_article: str, brand: str) -> str:
    cleaned = _remove_code_prefix(fragment, code_article)
    if not cleaned:
        return ""

    cleaned = cleaned.replace("ï€­", " ").replace("\uf0b7", " ").replace("â€¢", " ").replace("â—", " ").replace("â–ª", " ")
    cleaned = cleaned.replace("SAULER & PALAU", "SOLER & PALAU")
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")
    cleaned = re.sub(r"^(?:de\s+marque|marque|de\s+type|type)\s*[:\-]?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(
        r"^(?:fourniture(?:\s*,?\s*pose)?(?:\s+et\s+raccordement)?(?:\s+installation)?(?:\s+et\s+mise\s+en\s+service)?(?:\s+d['â€™]un[e]?)?\s*)",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(
        r"^(?:et\s+)?pose(?:\s+et\s+raccordement)?(?:\s+d['â€™]un[e]?)?\s*",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(
        r"^(?:et\s+)?mise\s+en\s+(?:oeuvre|Å“uvre|.{0,3}uvre)(?:\s+d['â€™]un[e]?)?\s*",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(
        r"^(?:est\s+mise\s+en\s+service\s+d['â€™]un[e]?\s*)",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(r"^d['â€™]?\s*un[e]?\s+", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^l['â€™]ensemble(?:\s+comprendra|\s+comporte)?\s*:?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^de\s+(?=(?:volet|gaine|cartouche|exutoire|ventilateur|grille|siphon|tube|tuyau|unite|unitÃ©|pompe|surpresseur)\b)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")

    normalized_cleaned = _normalize_search_text(cleaned)
    if any(hint in normalized_cleaned for hint in DOCUMENT_NOISE_HINTS):
        return ""
    if normalized_cleaned.startswith("contraintes d installation"):
        return ""
    if normalized_cleaned.startswith("selectionnee en fonction des besoins thermiques"):
        return ""

    if brand:
        brand_candidates = [candidate.strip() for candidate in brand.split("|") if candidate.strip()]
        if brand_candidates:
            brand_pattern = "|".join(re.escape(candidate) for candidate in brand_candidates)
            brand_sequence = rf"(?:{brand_pattern})(?:\s*(?:,|/|\bou\b|\bet\b)\s*(?:{brand_pattern}))*"
            cleaned = re.sub(
                rf",?\s*(?:de\s+)?marque\s+{brand_sequence}\s*(?:ou\s+(?:equivalent|Ã©quivalent|similaire))?",
                "",
                cleaned,
                flags=re.IGNORECASE,
            ).strip(" -:;,.")
            cleaned = re.sub(
                rf"\b{brand_sequence}\b\s*(?:ou\s+(?:equivalent|Ã©quivalent|similaire))?",
                "",
                cleaned,
                flags=re.IGNORECASE,
            ).strip(" -:;,.")

    cleaned = re.sub(r"\bou\s+(?:equivalent|Ã©quivalent|similaire)\b", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bmarque\b", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bde\s+(?=(?:comprenant|compose|avec)\b)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bde\.\s+(?=[A-ZÃ€-Ã¿])", "", cleaned)
    cleaned = re.sub(r"\bde\.?\s*$", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\b(?:ou|et)\b(?=\s*(?:,|/|\.|$))", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")
    return cleaned


def _extract_brand(text: str, product_parts: list[str] | None = None, chapter: str = "", subchapter: str = "") -> str:
    product_parts = product_parts or []
    cleaned_text = _normalize_ocr_punctuation(text)
    informative_lines = [
        _normalize_ocr_punctuation(line)
        for line in cleaned_text.splitlines()
        if _is_informative_line(line) and not any(hint in _normalize_search_text(line) for hint in DOCUMENT_NOISE_HINTS)
    ]

    scored_candidates: list[tuple[float, int, str]] = []
    for line_index, line in enumerate(informative_lines[:12]):
        explicit_groups = _extract_explicit_brand_groups(line)
        if explicit_groups:
            base_score = _score_brand_line(line, product_parts, chapter, subchapter) + 3.0
            if line_index < 4:
                base_score += 1.0
            for brands_group in explicit_groups:
                for candidate in _split_brand_candidates(brands_group):
                    scored_candidates.append((base_score, line_index, candidate))
            continue

        for pattern_index, pattern in enumerate(BRAND_PATTERNS):
            match = pattern.search(line)
            if not match:
                continue
            brands_group = match.groupdict().get("brands") or match.group(1)
            base_score = _score_brand_line(line, product_parts, chapter, subchapter) + max(0.0, 1.6 - (pattern_index * 0.2))
            if line_index < 4:
                base_score += 1.0
            for candidate in _split_brand_candidates(brands_group):
                scored_candidates.append((base_score, line_index, candidate))

    if not scored_candidates:
        normalized_text = _normalize_search_text(cleaned_text)
        fallback_patterns = (
            re.compile(
                r"(?:de marque|marque|de type|type|de chez|chez)\s+"
                r"(?P<brands>[a-z0-9+&./-]+(?:\s+[a-z0-9+&./-]+){0,3})"
                r"\s+ou\s+(?:equivalent|similaire)",
                re.IGNORECASE,
            ),
            re.compile(
                r"(?:de marque|marque|de type|type|de chez|chez)\s+"
                r"(?P<brands>[a-z0-9+&./-]+(?:\s+[a-z0-9+&./-]+){0,3})"
                r"\s+(?:comprenant|compose|avec)\b",
                re.IGNORECASE,
            ),
        )
        for pattern in fallback_patterns:
            match = pattern.search(normalized_text)
            if not match:
                continue
            brands_group = match.group("brands")
            for candidate in _split_brand_candidates(brands_group):
                scored_candidates.append((1.0, 99, candidate))
            if scored_candidates:
                break

    if not scored_candidates:
        return ""

    best_scores: dict[str, tuple[float, int]] = {}
    for score, line_index, candidate in scored_candidates:
        current = best_scores.get(candidate)
        if current is None or score > current[0] or (score == current[0] and line_index < current[1]):
            best_scores[candidate] = (score, line_index)

    uppercase_text = cleaned_text.upper()
    ranked = sorted(
        best_scores.items(),
        key=lambda item: (
            -item[1][0],
            item[1][1],
            uppercase_text.find(item[0]) if uppercase_text.find(item[0]) >= 0 else 10**6,
            -len(item[0]),
            item[0],
        ),
    )
    filtered_ranked: list[tuple[str, tuple[float, int]]] = []
    for candidate, meta in ranked:
        normalized_candidate = _normalize_search_text(candidate)
        is_subphrase = any(
            normalized_candidate
            and normalized_candidate != _normalize_search_text(other_candidate)
            and normalized_candidate in _normalize_search_text(other_candidate).split()
            for other_candidate, _ in ranked
        )
        if not is_subphrase:
            filtered_ranked.append((candidate, meta))

    ranked = filtered_ranked
    if ranked:
        best_score = ranked[0][1][0]
        ranked = [item for item in ranked if item[1][0] >= max(0.8, best_score - 2.5)]
    if not ranked or ranked[0][1][0] < 0.8:
        return ""
    return " | ".join(candidate for candidate, _ in ranked[:4])


def _sanitize_spec_fragment(fragment: str, code_article: str, brand: str) -> str:
    cleaned = _remove_code_prefix(fragment, code_article)
    if not cleaned:
        return ""

    cleaned = cleaned.replace("ï€­", " ").replace("\uf0b7", " ").replace("â€¢", " ").replace("â—", " ").replace("â–ª", " ")
    cleaned = cleaned.replace("SAULER & PALAU", "SOLER & PALAU")
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")
    cleaned = re.sub(r"^(?:de\s+marque|marque|de\s+type|type)\s*[:\-]?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(
        r"^(?:fourniture(?:\s*,?\s*pose)?(?:\s+et\s+raccordement)?(?:\s+installation)?(?:\s+et\s+mise\s+en\s+service)?(?:\s+d.?un(?:e)?)?\s*)",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(r"^(?:et\s+)?pose(?:\s+et\s+raccordement)?(?:\s+d.?un(?:e)?)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^(?:et\s+)?mise\s+en\s+(?:oeuvre|uvre)(?:\s+d.?un(?:e)?)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^(?:est\s+mise\s+en\s+service\s+d.?un(?:e)?\s*)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^d.?\s*un(?:e)?\s+", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^l.ensemble(?:\s+comprendra|\s+comporte)?\s*:?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^de\s+(?=(?:volet|gaine|cartouche|exutoire|ventilateur|grille|siphon|tube|tuyau|unite|unite|pompe|surpresseur)\b)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")

    normalized_cleaned = _normalize_search_text(cleaned)
    if any(hint in normalized_cleaned for hint in DOCUMENT_NOISE_HINTS):
        return ""
    if normalized_cleaned.startswith("contraintes d installation"):
        return ""
    if normalized_cleaned.startswith("selectionnee en fonction des besoins thermiques"):
        return ""

    if brand:
        brand_candidates = [candidate.strip() for candidate in brand.split("|") if candidate.strip()]
        if brand_candidates:
            brand_pattern = "|".join(re.escape(candidate) for candidate in brand_candidates)
            brand_sequence = rf"(?:{brand_pattern})(?:\s*(?:,|/|\bou\b|\bet\b)\s*(?:{brand_pattern}))*"
            cleaned = re.sub(
                rf",?\s*(?:de\s+)?marque\s+{brand_sequence}\s*(?:ou\s+(?:equivalent|equivalent|similaire))?",
                "",
                cleaned,
                flags=re.IGNORECASE,
            ).strip(" -:;,.")
            cleaned = re.sub(
                rf"\b{brand_sequence}\b\s*(?:ou\s+(?:equivalent|equivalent|similaire))?",
                "",
                cleaned,
                flags=re.IGNORECASE,
            ).strip(" -:;,.")

    cleaned = re.sub(r"\bou\s+(?:equivalent|equivalent|similaire)\b", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bmarque\b", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bde\s+(?=(?:comprenant|compose|avec)\b)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bde\.\s+(?=[A-Za-z])", "", cleaned)
    cleaned = re.sub(r"\bde\.?\s*$", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\b(?:ou|et)\b(?=\s*(?:,|/|\.|$))", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")
    return cleaned


def _sanitize_spec_fragment(fragment: str, code_article: str, brand: str) -> str:
    cleaned = _remove_code_prefix(fragment, code_article)
    if not cleaned:
        return ""

    cleaned = (
        cleaned.replace("ï€­", " ")
        .replace("\uf0b7", " ")
        .replace("â€¢", " ")
        .replace("â—", " ")
        .replace("â–ª", " ")
        .replace("•", " ")
        .replace("●", " ")
        .replace("▪", " ")
    )
    cleaned = cleaned.replace("SAULER & PALAU", "SOLER & PALAU")
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")
    cleaned = re.sub(r"^(?:de\s+marque|marque|de\s+type|type)\s*[:\-]?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(
        r"^(?:fourniture(?:\s*,?\s*pose)?(?:\s+et\s+raccordement)?(?:\s+installation)?(?:\s+et\s+mise\s+en\s+service)?(?:\s+d.?un(?:e)?)?\s*)",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(r"^(?:et\s+)?pose(?:\s+et\s+raccordement)?(?:\s+d.?un(?:e)?)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^(?:et\s+)?mise\s+en\s+(?:oeuvre|uvre)(?:\s+d.?un(?:e)?)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^(?:est\s+mise\s+en\s+service\s+d.?un(?:e)?\s*)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^d.?\s*un(?:e)?\s+", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^l.ensemble(?:\s+comprendra|\s+comporte)?\s*:?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^de\s+(?=(?:volet|gaine|cartouche|exutoire|ventilateur|grille|siphon|tube|tuyau|unite|pompe|surpresseur)\b)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")

    normalized_cleaned = _normalize_search_text(cleaned)
    if any(hint in normalized_cleaned for hint in DOCUMENT_NOISE_HINTS):
        return ""
    if normalized_cleaned.startswith("contraintes d installation"):
        return ""
    if normalized_cleaned.startswith("selectionnee en fonction des besoins thermiques"):
        return ""

    if brand:
        brand_candidates = [candidate.strip() for candidate in brand.split("|") if candidate.strip()]
        if brand_candidates:
            brand_pattern = "|".join(re.escape(candidate) for candidate in brand_candidates)
            brand_sequence = rf"(?:{brand_pattern})(?:\s*(?:,|/|\bou\b|\bet\b)\s*(?:{brand_pattern}))*"
            cleaned = re.sub(
                rf",?\s*(?:de\s+)?marque\s+{brand_sequence}\s*(?:ou\s+(?:equivalent|equivalent|similaire))?",
                "",
                cleaned,
                flags=re.IGNORECASE,
            ).strip(" -:;,.")
            cleaned = re.sub(
                rf"\b{brand_sequence}\b\s*(?:ou\s+(?:equivalent|equivalent|similaire))?",
                "",
                cleaned,
                flags=re.IGNORECASE,
            ).strip(" -:;,.")

    cleaned = re.sub(r"\bou\s+(?:equivalent|equivalent|similaire)\b", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bmarque\b", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bde\s+(?=(?:comprenant|compose|avec)\b)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bde\.\s+(?=[A-Za-z])", "", cleaned)
    cleaned = re.sub(r"\bde\.?\s*$", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\b(?:ou|et)\b(?=\s*(?:,|/|\.|$))", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"[\"«»]+\s*[\"«»]+", "", cleaned)
    cleaned = re.sub(r"\s+,", ",", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:;,.")
    return cleaned


def _pricing_status(record: dict[str, object]) -> str:
    numeric_values = []
    for key in ("prix_unitaire", "montant"):
        value = record.get(key)
        if value is None:
            continue
        text = str(value).strip()
        if not text:
            continue
        try:
            numeric_value = float(text)
        except (TypeError, ValueError):
            continue
        if math.isnan(numeric_value):
            continue
        numeric_values.append(numeric_value)

    if not numeric_values:
        return "a_renseigner"
    if all(abs(value) < 1e-9 for value in numeric_values):
        return "a_renseigner"
    return "renseigne"


def _extract_brand_and_specs_with_nvidia(excerpt: str, product_name: str, code_article: str, chapter: str, subchapter: str, max_retries: int = 3) -> tuple[str, str]:
    if not excerpt or not excerpt.strip():
        return "", ""

    import requests
    import json
    import time

    api_key = "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"
    invoke_url = "https://integrate.api.nvidia.com/v1/chat/completions"

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Accept": "application/json",
        "Content-Type": "application/json"
    }

    prompt = f"""Tu es un ingénieur technique expert en bâtiment et fluides (CVC/Plomberie/Électricité).
Analyse l'extrait (contexte) suivant tiré du descriptif technique d'un projet pour extraire de manière très précise la marque (fabricant) prescrite et les spécifications techniques associées à ce produit.

PRODUIT DU BORDEREAU :
- Code Article : {code_article}
- Désignation : {product_name}
- Localisation : {chapter} / {subchapter}

CONTEXTE EXTRACTIBLES :
{excerpt}

Consignes strictes d'extraction :
1. MARQUE (brand) :
   - Identifie le fabricant ou la marque du matériel (ex: Wilo, Daikin, Nicoll, S&P, Lowara, Danfoss, Ariston, Bayard, France Air, Aldes, Trox, etc.)
   - S'il y a plusieurs marques acceptées, sépare-les par un " | " (ex: Wilo | Lowara)
   - S'il y a "ou similaire", "ou équivalent", "ou équivalent", inclus la marque principale ET "ou similaire" ou "ou équivalent" (ex: Daikin | Mitsubishi | Toshiba ou similaire)
   - IMPORTANT: Ne supprime JAMAIS les mentions "ou similaire" ou "ou équivalent" - elles doivent apparaître TELLES QUELLES dans le résultat
   - Si aucune marque n'est mentionnée, laisse vide.

2. SPÉCIFICATIONS TECHNIQUES (specs) :
   - Extrait uniquement les caractéristiques physiques et techniques réelles clés (ex: diamètres DN50, pressions PN16, puissances 15kW, débits 12m3/h, type de matériau PVC/PPR/acier, exigences de température, coupe-feu, acoustique).
   - Synthétise en une seule ligne compacte et percutante.
   - Ne mets aucune phrase d'introduction ni de blabla.

Retourne uniquement un dictionnaire JSON valide sous le format exact :
{{
  "brand": "Nom(s) de marque OU VIDE",
  "specs": "Caractéristiques ou vide"
}}
N'écris rien d'autre que ce dictionnaire JSON. Le champ "brand" doit contenir les marques avec "ou similaire" ou "ou équivalent" si présent dans le texte."""

    payload = {
        "model": "mistralai/mistral-large-3-675b-instruct-2512",
        "messages": [
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.1,
        "max_tokens": 512,
        "stream": False
    }

    for attempt in range(max_retries):
        try:
            response = requests.post(invoke_url, headers=headers, json=payload, timeout=30)
            if response.status_code == 429:
                wait_time = (attempt + 1) * 2
                time.sleep(wait_time)
                continue
            response.raise_for_status()
            res_json = response.json()
            content = res_json["choices"][0]["message"]["content"].strip()

            if "```json" in content:
                content = content.split("```json")[1].split("```")[0].strip()
            elif "```" in content:
                content = content.split("```")[1].split("```")[0].strip()

            data = json.loads(content)
            return str(data.get("brand", "")).strip(), str(data.get("specs", "")).strip()
        except requests.exceptions.Timeout:
            if attempt < max_retries - 1:
                time.sleep(2)
                continue
            print(f"Timeout extraction NVIDIA pour {product_name}")
            return "", ""
        except Exception as e:
            if attempt < max_retries - 1:
                time.sleep(1)
                continue
            print(f"Erreur extraction NVIDIA pour {product_name}: {e}")
            return "", ""

    return "", ""


def enrich_products_dataframe(products_df: pd.DataFrame, descriptif_paths: list[Path], use_nvidia: bool = False) -> pd.DataFrame:
    corpus = build_pdf_corpus(descriptif_paths)
    repeated_signatures = _collect_repeated_line_signatures(corpus)
    rows: list[dict[str, object]] = []
    last_matched_page_by_chapter: dict[str, PdfPage] = {}

    # 1. Première passe : On fait correspondre les pages/blocs pour toutes les lignes
    records_to_process = []

    for record in products_df.to_dict(orient="records"):
        raw_product_name = str(record.get("product_name", "")).strip()
        product_name = _clean_display_product_name(raw_product_name)
        power_only_product = _is_power_only_product(product_name)
        code_article = str(record.get("code_article", "") or "").strip()
        local_subcode = _is_local_subcode(code_article)
        chapter = str(record.get("chapter", "") or "")
        subchapter = str(record.get("subchapter", "") or "")
        product_parts = _split_product_parts(product_name)
        best_block: MatchedBlock | None = None
        power_block: MatchedBlock | None = None
        is_noise = _is_financial_noise(code_article) or _is_financial_noise(product_name)

        if not is_noise and code_article and not local_subcode:
            best_block = _find_best_code_block(corpus, code_article, product_parts, chapter, subchapter)

        if not is_noise and power_only_product and code_article and not local_subcode:
            power_block = _find_power_table_block(corpus, code_article, chapter, subchapter)
            if power_block is not None and (best_block is None or power_block.score > best_block.score):
                best_block = power_block

        if best_block is None and not is_noise:
            candidate_page_sets: list[list[PdfPage]] = []
            chapter_anchor = last_matched_page_by_chapter.get(chapter)
            if chapter_anchor is not None:
                nearby_pages = [
                    page
                    for page in corpus
                    if page.pdf_path == chapter_anchor.pdf_path
                    and chapter_anchor.page_number - 1 <= page.page_number <= chapter_anchor.page_number + 3
                ]
                if nearby_pages:
                    candidate_page_sets.append(nearby_pages)

            if not local_subcode or not candidate_page_sets:
                candidate_page_sets.append(corpus)

            for candidate_pages in candidate_page_sets:
                local_best_page: PdfPage | None = None
                local_best_score = 0.0
                local_best_excerpt_lines: list[str] = []
                for page in candidate_pages:
                    score = _score_page(page, product_parts, chapter, subchapter)
                    excerpt_lines = _lines_around_match(page, product_parts, chapter, subchapter)
                    excerpt_text = "\n".join(excerpt_lines)
                    if re.search(r"de\s+marque|marque|ou\s+(?:equivalent|équivalent|similaire)|\btype\b|réf|ref\s*:", excerpt_text, re.IGNORECASE):
                        score += 2.0
                    if score > local_best_score:
                        local_best_score = score
                        local_best_page = page
                        local_best_excerpt_lines = excerpt_lines
                
                if local_best_page is not None and local_best_score >= 2.0:
                    best_page = local_best_page
                    best_score = local_best_score
                    best_excerpt_lines = local_best_excerpt_lines
                    best_block = MatchedBlock(page=best_page, lines=best_excerpt_lines, score=best_score, matched_code=code_article)
                    best_block = _extend_block_with_next_page(corpus, best_block, product_parts, chapter, subchapter)
                    break

        records_to_process.append({
            "record": record,
            "product_name": product_name,
            "code_article": code_article,
            "chapter": chapter,
            "subchapter": subchapter,
            "product_parts": product_parts,
            "best_block": best_block,
            "power_only_product": power_only_product,
        })

        if best_block and best_block.score >= (1.0 if power_only_product else 2.0) and chapter:
            last_matched_page_by_chapter[chapter] = best_block.page

    # 2. Deuxième passe : Extraction des données (Local ou NVIDIA en parallèle)
    tasks = []
    for item in records_to_process:
        best_block = item["best_block"]
        power_only_product = item["power_only_product"]
        minimum_score = 1.0 if power_only_product else 2.0
        
        excerpt = ""
        matched_lines = []
        if best_block and best_block.score >= minimum_score:
            matched_lines = _filter_noise_lines(best_block.lines, repeated_signatures, best_block.matched_code)
            excerpt = "\n".join(matched_lines[:12])
            
        tasks.append({
            "item": item,
            "excerpt": excerpt,
            "matched_lines": matched_lines,
        })

    # Si use_nvidia est activé, on prépare les requêtes en parallèle
    nvidia_results = {}
    if use_nvidia:
        from concurrent.futures import ThreadPoolExecutor
        
        def run_task(idx, excerpt, p_name, c_art, ch, sub_ch):
            if not excerpt or not excerpt.strip():
                return idx, ("", "")
            res = _extract_brand_and_specs_with_nvidia(excerpt, p_name, c_art, ch, sub_ch)
            return idx, res
            
        with ThreadPoolExecutor(max_workers=3) as executor:
            futures = []
            for idx, task in enumerate(tasks):
                item = task["item"]
                if task["excerpt"]:
                    futures.append(
                        executor.submit(
                            run_task,
                            idx,
                            task["excerpt"],
                            item["product_name"],
                            item["code_article"],
                            item["chapter"],
                            item["subchapter"]
                        )
                    )
            
            for fut in futures:
                try:
                    idx, res = fut.result()
                    nvidia_results[idx] = res
                except Exception as e:
                    print(f"ThreadPool Future exception: {e}")

    # 3. Troisième passe : Assemblage final des lignes du DataFrame
    for idx, task in enumerate(tasks):
        item = task["item"]
        best_block = item["best_block"]
        record = item["record"]
        matched_lines = task["matched_lines"]
        excerpt = task["excerpt"]
        
        brand = ""
        specs = ""
        page_number = None
        pdf_name = ""
        confidence = 0.0
        
        if best_block and best_block.score >= (1.0 if item["power_only_product"] else 2.0):
            page_number = best_block.page.page_number
            pdf_name = best_block.page.pdf_name
            confidence = min(round(best_block.score / 25.0, 3), 0.99)
            
            # Récupération des résultats NVIDIA
            if use_nvidia and idx in nvidia_results:
                n_brand, n_specs = nvidia_results[idx]
                brand = n_brand
                specs = n_specs
            
            # Fallback local s'il n'y a pas de résultat NVIDIA ou s'il est vide
            if not brand:
                matched_text = "\n".join(matched_lines)
                brand = _extract_brand(matched_text or best_block.page.text[:1600], item["product_parts"], item["chapter"], item["subchapter"])
                if not brand and _price_table_penalty(best_block.page) > 0.0:
                    brand = _extract_price_table_brand(corpus, best_block.page)
                    
            if not specs:
                specs = _extract_specs_compact(matched_lines, item["product_parts"], item["code_article"], brand)
                if not specs:
                    specs = _extract_specs_rich(matched_lines, item["product_parts"], brand)

        rows.append(
            {
                "sheet_name": record.get("sheet_name"),
                "row_index": record.get("row_index"),
                "chapter": item["chapter"],
                "subchapter": item["subchapter"],
                "code_article": item["code_article"],
                "product_name": item["product_name"],
                "unite": record.get("unite"),
                "quantite": record.get("quantite"),
                "matched_page": page_number,
                "matched_pdf": pdf_name,
                "brand": brand,
                "specs": specs,
                "matched_excerpt": excerpt,
                "match_confidence": confidence,
            }
        )

    return pd.DataFrame(rows)


def extract_products_from_descriptif(
    descriptif_paths: list[Path],
    workspace_root: Path,
    use_nvidia: bool = True,
) -> pd.DataFrame:
    from concurrent.futures import ThreadPoolExecutor
    import requests
    import json

    NVIDIA_API_KEY = "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"
    NVIDIA_BASE_URL = "https://integrate.api.nvidia.com/v1"

    corpus = build_pdf_corpus(descriptif_paths)

    page_texts = []
    for page in corpus:
        page_text = f"[Page {page.page_number} - {page.pdf_name}]\n" + "\n".join(page.lines)
        page_texts.append(page_text)

    batch_size = 8
    batches = [page_texts[i:i+batch_size] for i in range(0, len(page_texts), batch_size)]

    all_products: list[dict] = []

    system_prompt = """Tu es un expert en extraction de bordereau technique BTP.
Ton objectif est d'extraire les produits du descriptif technique en respectant STRICTEMENT le format ci-dessous.

STRUCTURE DES CHAPITRES (utilise ces noms EXACTS):
- A/ = ÉVACUATION EAUX PLUVIALES, EAUX USÉES ET EAUX VANNES (codes: A-01 à A-05 UNIQUEMENT)
- B/ = ALIMENTATION EAU FROIDE – EAU CHAUDE (codes: B-01 à B-09)
- C/ = APPAREILS SANITAIRES (POSE) (codes: C-01 à C-05)
- D/ = ACCESSOIRES (POSE) (codes: D-01 à D-04)
- E/ = PROTECTION INCENDIE (codes: E-01 à E-03)
- F/ = CLIMATISATION (codes: F-01 à F-17)
- G/ = VENTILATION MÉCANIQUE CONTRÔLÉE (codes: G-01 à G-07)
- H/ = DÉSENFUMAGE (codes: H-01 à H-06)

RÈGLES ABSOLUES - NE PAS CONTREDIRE:
1. Les codes A-01 à A-05 sont UNIQUEMENT: Tuyau PVC Évacuation, Tuyau PVC Pression, Siphon de Sol, Garde Grève, Station de Relevage
2. B-07 est "Équipement Bassin d'Eau" avec 8 sous-postes (1-8). Les sous-postes n'ont PAS de code_article, ils ont parent_code="B-07"
3. B-08 est "Attente Arrosage", B-09 est "Tube Polyéthylène Ligne Bleu"
4. Pour F-10 (Grille de Soufflage), le code F-10 est pour la version principale. Pour les sous-types a (Linéaire) et b (Carré), utilise parent_code="F-10", sub_code="a" ou "b"
5. Pour F-17 (Armoire Électrique et Câblage), sous-types: a1 (UI Clim Étages), a2 (Clim + Ventilation Terrasse), b (Câble et Chemin de Câble)
6. Pour H-03 (Arrêts Pompiers et Réarmements), le code H-03 est pour l'ensemble. Sous-types: a (Arrêts Pompier), b (Réarmement)
7. E-01 est "Extincteur Portatif" - les types (Eau pulv©risée 6L + CO2 2kg) sont des descriptions, pas des codes séparés
8. G-06 apparaît DEUX FOIS dans le document (Grille de Rejet ET Clapet Coupe-Feu 1H AutoCommandé) - Extraire les DEUX

NE PAS INVENTER:
- Ne crée PAS d'items avec les codes A-02, A-03, A-04, A-05 qui ne sont PAS "Tuyau PVC Pression", "Siphon de Sol", "Garde Grève", "Station de Relevage"
- Les items comme "Vanne de Vidange Ø40", "Crépine d'Aspiration Ø50", "Trappe Hermétique", "Trop Plein Ø100" ne sont PAS dans le document original avec ces codes

FORMAT JSON:
```json
[
  {"chapter": "A/", "subchapter": "ÉVACUATION EAUX PLUVIALES, EAUX USÉES ET EAUX VANNES", "code_article": "A-01", "product_name": "Tuyau PVC Évacuation", "unite": "ml"},
  {"chapter": "B/", "subchapter": "ALIMENTATION EAU FROIDE – EAU CHAUDE", "code_article": "B-07", "product_name": "Équipement Bassin d'Eau", "unite": "ensemble"},
  {"chapter": "B/", "subchapter": "ALIMENTATION EAU FROIDE – EAU CHAUDE", "parent_code": "B-07", "sub_code": "1", "product_name": "Poste de Filtration", "unite": "U"}
]
```

Retourne UNIQUEMENT un JSON array valide, sans texte avant ou après:"""

    def process_batch(batch_idx: int, batch_texts: list[str]) -> list[dict]:
        from openai import OpenAI

        combined_text = "\n\n===== PAGE SUIVANTE =====\n\n".join(batch_texts)

        prompt = f"{system_prompt}\n\nTEXTES DES PAGES:\n{combined_text}\n\nJSON array (uniquement):"

        try:
            client = OpenAI(base_url=NVIDIA_BASE_URL, api_key=NVIDIA_API_KEY)
            completion = client.chat.completions.create(
                model="qwen/qwen3-coder-480b-a35b-instruct",
                messages=[
                    {"role": "system", "content": "Tu es un expert technique BTP. Réponds uniquement en JSON array."},
                    {"role": "user", "content": prompt}
                ],
                temperature=0.1,
                max_tokens=4096,
                stream=False
            )

            if completion.choices and len(completion.choices) > 0:
                content = completion.choices[0].message.content or ""
                if "```json" in content:
                    content = content.split("```json")[1].split("```")[0].strip()
                elif "```" in content:
                    content = content.split("```")[1].split("```")[0].strip()

                content = content.strip()
                if content.startswith("[") and content.endswith("]"):
                    products = json.loads(content)
                    return products if isinstance(products, list) else []
        except Exception as e:
            print(f"Erreur batch {batch_idx}: {e}")

        return []

    with ThreadPoolExecutor(max_workers=4) as executor:
        futures = []
        for idx, batch in enumerate(batches):
            futures.append(executor.submit(process_batch, idx, batch))

        for fut in futures:
            try:
                products = fut.result()
                if products:
                    all_products.extend(products)
            except Exception as e:
                print(f"Erreur futur: {e}")

    rows = []
    seen_keys = set()
    for idx, prod in enumerate(all_products):
        if not isinstance(prod, dict):
            continue

        product_name = str(prod.get("product_name", "")).strip()
        if not product_name or len(product_name) < 3:
            continue

        chapter = str(prod.get("chapter", "")).strip()
        subchapter = str(prod.get("subchapter", "")).strip()
        code_article = str(prod.get("code_article", "")).strip()
        parent_code = str(prod.get("parent_code", "")).strip()
        sub_code = str(prod.get("sub_code", "")).strip()
        unite = str(prod.get("unite", "")).strip()
        quantite = str(prod.get("quantite", "")).strip()

        if any(kw in product_name.lower() for kw in ["prix", "total", "montant", "tva", "forfait"]):
            continue

        dedup_key = f"{chapter}|{subchapter}|{code_article}|{parent_code}|{sub_code}|{product_name}"
        if dedup_key in seen_keys:
            continue
        seen_keys.add(dedup_key)

        if unite == "" and code_article and not parent_code:
            if "E-01" in code_article or "extincteur" in product_name.lower():
                unite = "U"

        rows.append({
            "sheet_name": "descriptif",
            "row_index": idx,
            "chapter": chapter,
            "subchapter": subchapter,
            "code_article": code_article,
            "parent_code": parent_code,
            "sub_code": sub_code,
            "product_name": product_name,
            "unite": unite,
            "quantite": quantite,
            "prix_unitaire": "",
            "montant": "",
            "confidence": 0.7,
        })

    return pd.DataFrame(rows)


def export_enriched_excel(enriched_df: pd.DataFrame, destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(destination, engine="openpyxl") as writer:
        enriched_df.to_excel(writer, sheet_name="produits_enrichis", index=False)
        worksheet = writer.sheets["produits_enrichis"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        header_font = Font(bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_alignment = Alignment(vertical="top", wrap_text=True)

        preferred_widths = {
            "sheet_name": 14,
            "row_index": 10,
            "chapter": 24,
            "subchapter": 24,
            "code_article": 14,
            "product_name": 34,
            "unite": 10,
            "quantite": 12,
            "matched_page": 12,
            "brand": 22,
            "specs": 56,
            "matched_excerpt": 68,
            "match_confidence": 16,
        }

        headers = [cell.value for cell in worksheet[1]]
        width_map: dict[int, float] = {}
        for column_index, header in enumerate(headers, start=1):
            if not header:
                continue
            header_cell = worksheet.cell(row=1, column=column_index)
            header_cell.fill = header_fill
            header_cell.font = header_font
            header_cell.alignment = header_alignment

            target_width = preferred_widths.get(str(header), 18)
            width_map[column_index] = target_width
            worksheet.column_dimensions[get_column_letter(column_index)].width = target_width

        worksheet.sheet_view.showGridLines = True

        for row in worksheet.iter_rows(min_row=2):
            estimated_lines = 1
            for cell in row:
                cell.alignment = cell_alignment
                value = "" if cell.value is None else str(cell.value)
                if not value:
                    continue
                column_width = max(int(width_map.get(cell.column, 18)), 8)
                logical_lines = value.count("\n") + 1
                wrapped_lines = max(1, math.ceil(len(value) / max(column_width - 2, 8)))
                estimated_lines = max(estimated_lines, logical_lines, wrapped_lines)
            worksheet.row_dimensions[row[0].row].height = min(max(estimated_lines * 15, 18), 120)
    return destination
