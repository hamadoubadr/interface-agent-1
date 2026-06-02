from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
import xlrd

from .models import SheetData
from .normalizer import count_non_empty, is_empty


def _matrix_from_dataframe(frame: pd.DataFrame) -> list[list[Any]]:
    matrix: list[list[Any]] = []
    for row in frame.itertuples(index=False, name=None):
        matrix.append([None if pd.isna(value) else value for value in row])
    return matrix


def _sheet_dimensions(matrix: list[list[Any]]) -> tuple[int, int]:
    non_empty_rows = 0
    max_non_empty_cols = 0
    for row in matrix:
        row_count = count_non_empty(row)
        if row_count:
            non_empty_rows += 1
            max_non_empty_cols = max(max_non_empty_cols, row_count)
    return non_empty_rows, max_non_empty_cols


def _expand_merged_matrix(sheet: Worksheet) -> tuple[list[list[Any]], bool]:
    max_row = sheet.max_row or 0
    max_col = sheet.max_column or 0
    matrix: list[list[Any]] = [
        [sheet.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        for r in range(1, max_row + 1)
    ]
    expanded = False

    for merged_range in sheet.merged_cells.ranges:
        expanded = True
        min_col, min_row, max_col, max_row = merged_range.bounds
        anchor = sheet.cell(row=min_row, column=min_col).value
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                matrix[row - 1][col - 1] = anchor

    return matrix, expanded


def _read_xls_with_xlrd(workbook_path: Path) -> list[SheetData]:
    workbook = xlrd.open_workbook(workbook_path, formatting_info=True)
    sheets: list[SheetData] = []

    for sheet in workbook.sheets():
        matrix: list[list[Any]] = [
            [sheet.cell_value(row, col) for col in range(sheet.ncols)]
            for row in range(sheet.nrows)
        ]
        merged_expanded = False
        for row_start, row_end, col_start, col_end in getattr(sheet, "merged_cells", []):
            merged_expanded = True
            anchor = sheet.cell_value(row_start, col_start)
            for row in range(row_start, row_end):
                for col in range(col_start, col_end):
                    matrix[row][col] = anchor

        non_empty_rows, non_empty_cols = _sheet_dimensions(matrix)
        sheets.append(
            SheetData(
                name=sheet.name,
                matrix=matrix,
                source="xlrd",
                warnings=[],
                merged_cells_expanded=merged_expanded,
                formula_density=0.0,
                non_empty_rows=non_empty_rows,
                non_empty_cols=non_empty_cols,
            )
        )
    return sheets


def read_workbook(path: str | Path) -> list[SheetData]:
    workbook_path = Path(path)
    suffix = workbook_path.suffix.lower()
    if suffix == ".xls":
        return _read_xls_with_xlrd(workbook_path)

    excel_file = pd.ExcelFile(workbook_path, engine="openpyxl")
    workbook = load_workbook(workbook_path, data_only=True)

    sheets: list[SheetData] = []
    for sheet_name in excel_file.sheet_names:
        warnings: list[str] = []
        pandas_matrix: list[list[Any]] = []
        openpyxl_matrix: list[list[Any]] = []
        merged_expanded = False

        try:
            frame = excel_file.parse(sheet_name=sheet_name, header=None, dtype=object)
            pandas_matrix = _matrix_from_dataframe(frame)
        except Exception as exc:  # pragma: no cover - defensive branch
            warnings.append(f"Lecture pandas impossible: {exc}")

        worksheet = workbook[sheet_name]
        try:
            openpyxl_matrix, merged_expanded = _expand_merged_matrix(worksheet)
        except Exception as exc:  # pragma: no cover - defensive branch
            warnings.append(f"Lecture openpyxl impossible: {exc}")

        pandas_non_empty = sum(1 for row in pandas_matrix for cell in row if not is_empty(cell))
        openpyxl_non_empty = sum(1 for row in openpyxl_matrix for cell in row if not is_empty(cell))

        # Prefer the openpyxl view when merged cells hide data in the pandas parse.
        if merged_expanded or openpyxl_non_empty > pandas_non_empty * 1.05:
            matrix = openpyxl_matrix or pandas_matrix
            source = "openpyxl"
        else:
            matrix = pandas_matrix or openpyxl_matrix
            source = "pandas"

        non_empty_rows, non_empty_cols = _sheet_dimensions(matrix)
        sheets.append(
            SheetData(
                name=sheet_name,
                matrix=matrix,
                source=source,
                warnings=warnings,
                merged_cells_expanded=merged_expanded,
                formula_density=0.0,
                non_empty_rows=non_empty_rows,
                non_empty_cols=non_empty_cols,
            )
        )

    return sheets
